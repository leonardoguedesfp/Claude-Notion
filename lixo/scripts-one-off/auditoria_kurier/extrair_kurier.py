# -*- coding: utf-8 -*-
r"""
Extrai publicações dos e-mails Kurier de 2026 e gera planilha de auditoria.

Uso:
    .\.venv\Scripts\python.exe scripts\auditoria_kurier\extrair_kurier.py

Saída: Auditoria_Kurier_2026.xlsx na mesma pasta dos e-mails (fora do repo).

Lê todos os .eml em PASTA_EMAILS recursivamente, filtra os de 2026 (header Date),
parseia o HTML (formato DJEN com campos rotulados, formato STF "Novo" com blocos
[INFORMAÇÕES]/[PARTES]/[CONTEÚDO], e formatos livres tipo TJMG/TRF-1/TRT-10 alt),
deduplica por hash blake2b(processo+data+texto), e grava 3 abas (Publicacoes_Kurier,
Resumo, Log).

Dependências: beautifulsoup4, openpyxl (instaladas no venv).
"""

from __future__ import annotations

import email
import hashlib
import os
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from email import policy
from email.utils import parsedate_to_datetime
from pathlib import Path

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# -----------------------------------------------------------------------------
# Configuração
# -----------------------------------------------------------------------------

PASTA_EMAILS = Path(
    r"C:\Users\LeonardoGuedesdaFons\OneDrive - RICARDO PASSOS ADVOCACIA"
    r"\Área de Trabalho\E-mails exemplo"
)
ARQUIVO_SAIDA = PASTA_EMAILS / "Auditoria_Kurier_2026.xlsx"
ANO_ALVO = 2026

# 6 advogados do escritório (nome, OAB, UF). Nome em UPPER sem acento para comparar.
ESCRITORIO: list[tuple[str, str, str]] = [
    ("RICARDO LUIZ RODRIGUES DA FONSECA PASSOS", "15523", "DF"),
    ("LEONARDO GUEDES DA FONSECA PASSOS", "36129", "DF"),
    ("VITOR GUEDES DA FONSECA PASSOS", "48468", "DF"),
    ("CECILIA MARIA LAPETINA CHIARATTO", "20120", "DF"),
    ("SAMANTHA LAIS SOARES MICKIEVICZ", "38809", "DF"),
    ("DEBORAH NASCIMENTO DE CASTRO", "75799", "DF"),
]

# Caracteres de controle ilegais em XLSX (regra do projeto principal).
RE_CTRL_ILEGAL = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F␀-␿]")

COLUNAS = [
    "arquivo_origem",
    "data_email",
    "assunto_email",
    "advogado_destinatario_principal",
    "tribunal_kurier",
    "numero_processo",
    "data_disponibilizacao",
    "tipo_comunicacao",
    "meio",
    "orgao",
    "advogados_intimados",
    "oabs_escritorio_intimadas",
    "texto",
    "hash_kurier",
]


# -----------------------------------------------------------------------------
# Estruturas
# -----------------------------------------------------------------------------


@dataclass
class Linha:
    arquivo_origem: str = ""
    data_email: str = ""
    assunto_email: str = ""
    advogado_destinatario_principal: str = ""
    tribunal_kurier: str = ""
    numero_processo: str = ""
    data_disponibilizacao: str = ""
    tipo_comunicacao: str = ""
    meio: str = ""
    orgao: str = ""
    advogados_intimados: str = ""
    oabs_escritorio_intimadas: str = ""
    texto: str = ""
    hash_kurier: str = ""

    def como_tupla(self) -> tuple:
        return tuple(getattr(self, c) for c in COLUNAS)


@dataclass
class LogEntry:
    arquivo: str
    nivel: str  # INFO / WARN / ERROR
    mensagem: str


@dataclass
class Estado:
    linhas: list[Linha] = field(default_factory=list)
    log: list[LogEntry] = field(default_factory=list)
    arquivos_2026: int = 0
    arquivos_ignorados_ano: int = 0
    arquivos_com_erro: int = 0
    duplicatas_descartadas: int = 0
    hashes_vistos: set[str] = field(default_factory=set)

    def add_log(self, arquivo: str, nivel: str, msg: str) -> None:
        self.log.append(LogEntry(arquivo=arquivo, nivel=nivel, mensagem=msg))


# -----------------------------------------------------------------------------
# Utilidades
# -----------------------------------------------------------------------------


def sem_acento(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def sanitizar(s: str | None) -> str:
    if s is None:
        return ""
    if not isinstance(s, str):
        s = str(s)
    return RE_CTRL_ILEGAL.sub("", s)


def matches_oab_escritorio(texto: str, oab: str, uf: str) -> bool:
    """Verifica se OAB+UF aparece em qualquer formato comum no texto."""
    # Variantes do número (com/sem zeros à esquerda, comprimentos 4..7).
    variantes = {oab.lstrip("0"), oab.zfill(5), oab.zfill(6), oab.zfill(7)}
    for o in variantes:
        if not o:
            continue
        # Frente: 15523/DF, 15523-DF, 15523 DF
        if re.search(rf"(?<!\d){re.escape(o)}\s*[/\-\s]\s*{uf}\b", texto):
            return True
        # Reverso: DF015523, DF/15523, DF-15523, DF 15523
        if re.search(rf"\b{uf}\s*[/\-\s]?\s*{re.escape(o)}(?!\d)", texto):
            return True
        # Bloco DJEN: "Nº OAB:15523" + "UF:DF" em linhas próximas
        if re.search(
            rf"OAB[:\s]*{re.escape(o)}\b[\s\S]{{0,40}}UF[:\s]*{uf}\b",
            texto,
        ):
            return True
    return False


def render_oabs_escritorio(texto: str) -> str:
    """Retorna 'NOME (OAB/UF); NOME (OAB/UF)' dos advogados do escritório intimados."""
    encontrados = []
    for nome, oab, uf in ESCRITORIO:
        if matches_oab_escritorio(texto, oab, uf):
            encontrados.append(f"{nome} ({oab}/{uf})")
    return "; ".join(encontrados)


def hash_publicacao(numero_processo: str, data_disp: str, texto: str) -> str:
    base = f"{numero_processo}||{data_disp}||{texto}".encode("utf-8", errors="replace")
    return hashlib.blake2b(base, digest_size=16).hexdigest()


# -----------------------------------------------------------------------------
# Parsing de e-mail
# -----------------------------------------------------------------------------


_RE_PARA_DIA = re.compile(r"Para\s+(.+?)\s+do\s+Dia\s*:", re.IGNORECASE | re.DOTALL)
_RE_TRIBUNAL_SUBJECT = re.compile(
    r"Pesquisas\s+do\s+Di[áa]rio\s*[:\-]?\s*([^P\n]+?)\s+Para\s",
    re.IGNORECASE,
)


def extrair_advogado_destinatario(subject: str) -> str:
    m = _RE_PARA_DIA.search(subject)
    if not m:
        return ""
    nome = re.sub(r"\s+", " ", m.group(1)).strip()
    return nome


def tribunal_do_subject(subject: str) -> str:
    m = _RE_TRIBUNAL_SUBJECT.search(subject)
    if not m:
        return ""
    raw = re.sub(r"\s+", " ", m.group(1)).strip()
    return raw


# -----------------------------------------------------------------------------
# Parsing de publicação dentro do TextoProcessoOriginal
# -----------------------------------------------------------------------------

# Labels DJEN. As variantes cobrem flexões de acento eventuais.
_LABELS = {
    "PROCESSO": [r"PROCESSO\s*:"],
    "ORGAO": [r"ORG[ÃA]O\s*:", r"[ÓO]RG[ÃA]O\s*:"],
    "DATA_DISP": [r"DATA\s+DE\s+DISPONIBILIZA[ÇC][ÃA]O\s*:"],
    "TIPO_COM": [r"TIPO\s+DE\s+COMUNICA[ÇC][ÃA]O\s*:"],
    "MEIO": [r"MEIO\s*:"],
    "TRIBUNAL": [r"TRIBUNAL\s*:"],
    "TEXTO": [r"TEXTO\s*:"],
}

_RE_LABEL_FIM = re.compile(
    r"\n(?:PARTES|ADVOGADOS|INTIMADO\(S\))\b", re.IGNORECASE
)

# Padrões diversos de OAB para o "advogados_intimados" em formatos não-DJEN.
_RE_OAB_INLINE = re.compile(
    r"""
    (?P<nome>[A-ZÁÉÍÓÚÂÊÔÃÕÇ][A-ZÁÉÍÓÚÂÊÔÃÕÇ \.\-']{4,80}?)
    \s*\(?\s*OAB[:\s/]*
    (?:(?P<uf1>[A-Z]{2})\s*[\-/]?\s*(?P<n1>\d{3,7})|(?P<n2>\d{3,7})\s*[\-/]\s*(?P<uf2>[A-Z]{2}))
    \s*\)?
    """,
    re.VERBOSE,
)

_RE_CNJ = re.compile(r"\b\d{7}-\d{2}\.\d{4}\.\d{1,2}\.\d{2}\.\d{4}\b")

# Processo no formato STF Novo: "Processo = RE 1555538" / "Processo = RE 1555538 Mérito"
_RE_PROCESSO_LIVRE = re.compile(r"Processo\s*[=:]\s*([^\n]+)")

# Sufixo dos IDs (Gmail Forward prefixa com "m_-XXXX")
_RE_ID_TEXTO = re.compile(r"TextoProcessoOriginal$")
_RE_ID_PROC = re.compile(r"IdProcesso$")

# STF Novo Formato: nome em uma linha, '"OAB N/UF"' na linha seguinte.
_RE_STF_ADV = re.compile(
    r'^([A-Za-zÁÉÍÓÚÂÊÔÃÕÇáéíóúâêôãõç][A-Za-zÁÉÍÓÚÂÊÔÃÕÇáéíóúâêôãõç \.\-\']{4,80})\s*\n'
    r'\s*"?\s*OAB[s\(]*\s*[:\-]?\s*\(?([^"\n]+?)\)?\s*"?\s*$',
    re.MULTILINE,
)


def texto_de_span(span) -> str:
    """Converte <br> em \\n e retorna texto sem outras tags."""
    for br in span.find_all(["br"]):
        br.replace_with("\n")
    raw = span.get_text()
    # Normaliza espaços em torno de quebras
    raw = re.sub(r"[ \t]+\n", "\n", raw)
    raw = re.sub(r"\n[ \t]+", "\n", raw)
    return raw.strip()


def buscar_label(texto: str, padroes: list[str]) -> tuple[int, int] | None:
    for pat in padroes:
        m = re.search(pat, texto)
        if m:
            return m.start(), m.end()
    return None


def extrair_campo_simples(texto: str, padroes_label: list[str]) -> str:
    """Extrai valor de label simples (uma linha)."""
    pos = buscar_label(texto, padroes_label)
    if not pos:
        return ""
    fim_label = pos[1]
    # Valor termina na próxima quebra de linha
    nl = texto.find("\n", fim_label)
    valor = texto[fim_label:nl] if nl != -1 else texto[fim_label:]
    return valor.strip()


def extrair_texto_djen(texto: str) -> str:
    """Extrai valor do label TEXTO até o próximo PARTES/ADVOGADOS/fim."""
    pos = buscar_label(texto, _LABELS["TEXTO"])
    if not pos:
        return ""
    fim_label = pos[1]
    resto = texto[fim_label:]
    m_fim = _RE_LABEL_FIM.search(resto)
    if m_fim:
        return resto[: m_fim.start()].strip()
    return resto.strip()


def extrair_advogados_djen(texto: str) -> list[tuple[str, str, str]]:
    """Bloco ADVOGADOS\\nNOME:..\\nNº OAB:..\\nUF:.. (DJEN puro)."""
    m = re.search(r"\nADVOGADOS\s*\n", texto)
    if not m:
        return []
    bloco = texto[m.end() :]
    advs: list[tuple[str, str, str]] = []
    # Match triplas NOME/OAB/UF, com tolerância para espaços.
    re_tripla = re.compile(
        r"NOME\s*:\s*(?P<nome>.+?)\s*\n\s*N[ºo°]?\s*OAB\s*:\s*(?P<oab>\d+)\s*\n\s*UF\s*:\s*(?P<uf>[A-Z]{2})",
        re.IGNORECASE,
    )
    for tm in re_tripla.finditer(bloco):
        nome = re.sub(r"\s+", " ", tm.group("nome")).strip()
        # Remove asteriscos/HTML residual eventual
        nome = re.sub(r"^[\W_]+|[\W_]+$", "", nome).strip()
        if nome:
            advs.append((nome, tm.group("oab"), tm.group("uf").upper()))
    return advs


def extrair_advogados_inline(texto: str) -> list[tuple[str, str, str]]:
    """Captura nomes seguidos de OAB em formatos livres (TJMG, TRT-10 alt, etc.)."""
    advs: list[tuple[str, str, str]] = []
    for m in _RE_OAB_INLINE.finditer(texto):
        nome = re.sub(r"\s+", " ", m.group("nome")).strip(" -.")
        oab = m.group("n1") or m.group("n2") or ""
        uf = (m.group("uf1") or m.group("uf2") or "").upper()
        # Filtra falsos positivos óbvios (palavras curtas ou termos não-nome)
        if len(nome) < 5 or "OAB" in nome.upper():
            continue
        advs.append((nome, oab, uf))
    return advs


def extrair_advogados_stf(texto: str) -> list[tuple[str, str, str]]:
    """STF Novo Formato: nome (mixed case) em linha, '"OAB N/UF"' na seguinte."""
    advs: list[tuple[str, str, str]] = []
    # Pattern: NOME\n"OAB 12345/UF" ou "OABs (123/UF, 456/UF)" — extrai todos OAB/UF do quoted str
    pat = re.compile(
        r'^([A-Za-zÁÉÍÓÚÂÊÔÃÕÇáéíóúâêôãõç][A-Za-zÁÉÍÓÚÂÊÔÃÕÇáéíóúâêôãõç \.\-\']{4,80})\s*\n'
        r'\s*"OABs?\s*\(?\s*([^"]+?)\s*\)?\s*"',
        re.MULTILINE,
    )
    for m in pat.finditer(texto):
        nome = re.sub(r"\s+", " ", m.group(1)).strip(" -.")
        bloco_oabs = m.group(2)
        # Extrai todos os pares NUMERO/UF do bloco (ex: "68823/GO, 425551/SP, 16785/DF")
        pares = re.findall(r"(\d{3,7})\s*/\s*([A-Z]{2})", bloco_oabs)
        if pares:
            for oab, uf in pares:
                advs.append((nome, oab, uf.upper()))
        else:
            advs.append((nome, "", ""))
    return advs


def parse_publicacao(span_texto, span_id_proc) -> dict:
    """Extrai campos de uma publicação. Retorna dict pronto para Linha."""
    texto = texto_de_span(span_texto)
    res: dict[str, str] = {
        "numero_processo": "",
        "data_disponibilizacao": "",
        "tipo_comunicacao": "",
        "meio": "",
        "orgao": "",
        "tribunal_kurier": "",
        "advogados_intimados": "",
        "texto": texto,
    }

    # Tenta formato DJEN primeiro (labels canônicos).
    if buscar_label(texto, _LABELS["PROCESSO"]):
        res["numero_processo"] = extrair_campo_simples(texto, _LABELS["PROCESSO"])
        res["orgao"] = extrair_campo_simples(texto, _LABELS["ORGAO"])
        res["data_disponibilizacao"] = extrair_campo_simples(texto, _LABELS["DATA_DISP"])
        res["tipo_comunicacao"] = extrair_campo_simples(texto, _LABELS["TIPO_COM"])
        res["meio"] = extrair_campo_simples(texto, _LABELS["MEIO"])
        res["tribunal_kurier"] = extrair_campo_simples(texto, _LABELS["TRIBUNAL"])
        texto_publ = extrair_texto_djen(texto)
        if texto_publ:
            res["texto"] = texto_publ

        advs = extrair_advogados_djen(texto)
        if not advs:
            # Alguns DJEN não têm bloco ADVOGADOS — cair no inline
            advs = extrair_advogados_inline(texto)
        res["advogados_intimados"] = "; ".join(
            f"{n} (OAB {o}/{u})" if o and u else n for n, o, u in advs
        )
    else:
        # Formato livre (STF Novo Formato, TJMG, TRF-1, TRT-10 alt, etc.)
        # Processo: tenta CNJ primeiro, depois "Processo = SIGLA NUMERO" (STF Novo).
        m = _RE_CNJ.search(texto)
        if m:
            res["numero_processo"] = m.group(0)
        else:
            m2 = _RE_PROCESSO_LIVRE.search(texto)
            if m2:
                res["numero_processo"] = m2.group(1).strip()

        # Data publicação no STF Novo: "Publicação = YYYY-MM-DD"
        m_dp = re.search(r"Publica[çc][ãa]o\s*=\s*(\d{4}-\d{2}-\d{2})", texto)
        if m_dp:
            res["data_disponibilizacao"] = m_dp.group(1)

        # Advogados: tenta inline (TJMG/TRT-10 alt/TRF-1) e formato STF (multilinha).
        advs = extrair_advogados_inline(texto)
        if not advs:
            advs = extrair_advogados_stf(texto)
        res["advogados_intimados"] = "; ".join(
            f"{n} (OAB {o}/{u})" if o and u else (f"{n} ({o or u})" if (o or u) else n)
            for n, o, u in advs
        )

    # OABs do escritório intimadas (sempre roda no texto cru).
    res["oabs_escritorio_intimadas"] = render_oabs_escritorio(texto)

    return res


# -----------------------------------------------------------------------------
# Loop principal
# -----------------------------------------------------------------------------


def processar_arquivo(path: Path, est: Estado) -> None:
    nome = path.name
    try:
        with open(path, "rb") as f:
            msg = email.message_from_binary_file(f, policy=policy.default)
    except Exception as exc:
        est.arquivos_com_erro += 1
        est.add_log(nome, "ERROR", f"Falha ao abrir/parsear .eml: {exc}")
        return

    # Date header → datetime → ano
    raw_date = msg.get("Date") or ""
    dt = None
    if raw_date:
        try:
            dt = parsedate_to_datetime(raw_date)
        except Exception as exc:
            est.add_log(nome, "WARN", f"Date header não parseável ({raw_date!r}): {exc}")

    if dt is None or dt.year != ANO_ALVO:
        est.arquivos_ignorados_ano += 1
        return

    est.arquivos_2026 += 1
    data_email_iso = dt.strftime("%Y-%m-%d")

    subject = msg.get("Subject") or ""
    advogado_dest = extrair_advogado_destinatario(subject)
    tribunal_subject = tribunal_do_subject(subject)

    # Conteúdo HTML
    html: str = ""
    try:
        if msg.is_multipart():
            for part in msg.walk():
                if part.get_content_type() == "text/html":
                    html = part.get_content()
                    break
        else:
            html = msg.get_content() or ""
    except Exception as exc:
        est.arquivos_com_erro += 1
        est.add_log(nome, "ERROR", f"Falha ao extrair HTML: {exc}")
        return

    if not html:
        est.add_log(nome, "WARN", "Sem conteúdo HTML encontrado")
        return

    soup = BeautifulSoup(html, "html.parser")
    # Gmail Forward prefixa IDs com "m_-XXXX". Match por sufixo.
    spans_texto = soup.find_all("span", id=_RE_ID_TEXTO)
    spans_id = soup.find_all("span", id=_RE_ID_PROC)

    if not spans_texto:
        est.add_log(nome, "WARN", "Nenhuma publicação (TextoProcessoOriginal) encontrada")
        return

    pasta_tribunal = path.parent.name  # ex: 'TRT-10', usado como fallback

    for idx, span in enumerate(spans_texto):
        span_id = spans_id[idx] if idx < len(spans_id) else None
        try:
            campos = parse_publicacao(span, span_id)
        except Exception as exc:
            est.add_log(
                nome,
                "ERROR",
                f"Falha parse publicação #{idx+1}: {exc}",
            )
            continue

        # Fallbacks
        if not campos["tribunal_kurier"]:
            # Pasta é canônica (STF, TRT-10, TJMG…); subject só ajuda se pasta vazia.
            campos["tribunal_kurier"] = pasta_tribunal or tribunal_subject
        if not campos["data_disponibilizacao"]:
            campos["data_disponibilizacao"] = data_email_iso

        # Sanitização e build
        texto_san = sanitizar(campos["texto"])
        h = hash_publicacao(
            sanitizar(campos["numero_processo"]),
            sanitizar(campos["data_disponibilizacao"]),
            texto_san,
        )

        if h in est.hashes_vistos:
            est.duplicatas_descartadas += 1
            continue
        est.hashes_vistos.add(h)

        if not campos["numero_processo"]:
            est.add_log(
                nome,
                "WARN",
                f"Publicação #{idx+1} sem número de processo extraído",
            )

        linha = Linha(
            arquivo_origem=sanitizar(nome),
            data_email=data_email_iso,
            assunto_email=sanitizar(subject),
            advogado_destinatario_principal=sanitizar(advogado_dest),
            tribunal_kurier=sanitizar(campos["tribunal_kurier"]),
            numero_processo=sanitizar(campos["numero_processo"]),
            data_disponibilizacao=sanitizar(campos["data_disponibilizacao"]),
            tipo_comunicacao=sanitizar(campos["tipo_comunicacao"]),
            meio=sanitizar(campos["meio"]),
            orgao=sanitizar(campos["orgao"]),
            advogados_intimados=sanitizar(campos["advogados_intimados"]),
            oabs_escritorio_intimadas=sanitizar(campos["oabs_escritorio_intimadas"]),
            texto=texto_san,
            hash_kurier=h,
        )
        est.linhas.append(linha)


# -----------------------------------------------------------------------------
# Geração do Excel
# -----------------------------------------------------------------------------


def construir_excel(est: Estado, destino: Path) -> None:
    wb = Workbook()

    # ---------- Aba Publicacoes_Kurier ----------
    ws = wb.active
    ws.title = "Publicacoes_Kurier"
    ws.append(COLUNAS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"

    for linha in est.linhas:
        ws.append(linha.como_tupla())

    # Larguras
    larguras = {
        "arquivo_origem": 60,
        "data_email": 12,
        "assunto_email": 70,
        "advogado_destinatario_principal": 36,
        "tribunal_kurier": 12,
        "numero_processo": 28,
        "data_disponibilizacao": 14,
        "tipo_comunicacao": 18,
        "meio": 36,
        "orgao": 40,
        "advogados_intimados": 60,
        "oabs_escritorio_intimadas": 50,
        "texto": 80,
        "hash_kurier": 34,
    }
    for i, c in enumerate(COLUNAS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = larguras.get(c, 20)

    # ---------- Aba Resumo ----------
    ws_res = wb.create_sheet("Resumo")
    ws_res.append(["Métrica", "Valor"])
    ws_res["A1"].font = Font(bold=True)
    ws_res["B1"].font = Font(bold=True)

    ws_res.append(["E-mails de 2026 processados", est.arquivos_2026])
    ws_res.append(["E-mails ignorados (ano != 2026)", est.arquivos_ignorados_ano])
    ws_res.append(["E-mails com erro de parsing", est.arquivos_com_erro])
    ws_res.append(["Publicações extraídas (após dedup)", len(est.linhas)])
    ws_res.append(["Duplicatas descartadas (hash)", est.duplicatas_descartadas])
    ws_res.append([])

    # Distribuição por mês
    ws_res.append(["Distribuição por mês (data_email)"])
    ws_res[ws_res.max_row][0].font = Font(bold=True)
    cont_mes: Counter[str] = Counter()
    for l in est.linhas:
        if l.data_email and len(l.data_email) >= 7:
            cont_mes[l.data_email[:7]] += 1
    for k in sorted(cont_mes):
        ws_res.append([k, cont_mes[k]])
    ws_res.append([])

    # Distribuição por tribunal
    ws_res.append(["Distribuição por tribunal_kurier"])
    ws_res[ws_res.max_row][0].font = Font(bold=True)
    cont_trib: Counter[str] = Counter(l.tribunal_kurier for l in est.linhas)
    for k, v in sorted(cont_trib.items(), key=lambda x: (-x[1], x[0])):
        ws_res.append([k or "(vazio)", v])
    ws_res.append([])

    # Distribuição por advogado destinatário principal
    ws_res.append(["Distribuição por advogado_destinatario_principal"])
    ws_res[ws_res.max_row][0].font = Font(bold=True)
    cont_dest: Counter[str] = Counter(l.advogado_destinatario_principal for l in est.linhas)
    for k, v in sorted(cont_dest.items(), key=lambda x: (-x[1], x[0])):
        ws_res.append([k or "(vazio)", v])
    ws_res.append([])

    # Quantidade de publicações com cada um dos 6 advogados do escritório
    ws_res.append(["Publicações com OAB do escritório intimada"])
    ws_res[ws_res.max_row][0].font = Font(bold=True)
    for nome, oab, uf in ESCRITORIO:
        chave = f"{nome} ({oab}/{uf})"
        c = sum(1 for l in est.linhas if chave in l.oabs_escritorio_intimadas)
        ws_res.append([chave, c])

    ws_res.column_dimensions["A"].width = 60
    ws_res.column_dimensions["B"].width = 14

    # ---------- Aba Log ----------
    ws_log = wb.create_sheet("Log")
    ws_log.append(["arquivo", "nivel", "mensagem"])
    for cell in ws_log[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    for entry in est.log:
        ws_log.append([sanitizar(entry.arquivo), entry.nivel, sanitizar(entry.mensagem)])
    ws_log.column_dimensions["A"].width = 60
    ws_log.column_dimensions["B"].width = 8
    ws_log.column_dimensions["C"].width = 90

    wb.save(destino)


# -----------------------------------------------------------------------------
# Entry point
# -----------------------------------------------------------------------------


def main() -> int:
    if not PASTA_EMAILS.exists():
        print(f"ERRO: pasta dos e-mails não encontrada: {PASTA_EMAILS}", file=sys.stderr)
        return 2
    arquivos = sorted(PASTA_EMAILS.rglob("*.eml"))
    if not arquivos:
        print(f"ERRO: pasta vazia (nenhum .eml): {PASTA_EMAILS}", file=sys.stderr)
        return 3

    print(f"Pasta: {PASTA_EMAILS}")
    print(f"Total de .eml encontrados: {len(arquivos)}")
    print(f"Filtro de ano: {ANO_ALVO}")
    print()

    est = Estado()
    n = len(arquivos)
    for i, p in enumerate(arquivos, start=1):
        if i % 50 == 0 or i == n:
            print(f"  [{i:>4}/{n}] processando…")
        processar_arquivo(p, est)

    print()
    print(f"E-mails de {ANO_ALVO} processados : {est.arquivos_2026}")
    print(f"E-mails ignorados (outro ano)    : {est.arquivos_ignorados_ano}")
    print(f"E-mails com erro                 : {est.arquivos_com_erro}")
    print(f"Publicações extraídas (final)    : {len(est.linhas)}")
    print(f"Duplicatas descartadas           : {est.duplicatas_descartadas}")
    print(f"Entradas de log                  : {len(est.log)}")
    print()
    print(f"Gravando: {ARQUIVO_SAIDA}")
    construir_excel(est, ARQUIVO_SAIDA)
    print("OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
