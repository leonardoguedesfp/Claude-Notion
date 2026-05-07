"""Análise dos disparos de alertas — Round 7 (2026-05-04).

Cruza CSVs exportados de 📬 Publicações × ⚖️ Processos para diagnosticar
caso a caso o motivo de cada alerta disparar. Saída:

- ``logs/analise_alertas_round_7.xlsx`` — uma aba por alerta com >50 disparos,
  uma aba ``Outros`` para os ≤50, mais aba ``Resumo`` com navegação.

NÃO modifica produto. Análise pura.

Uso (do main repo):
    python scripts/analise/analise_alertas_round_7.py [--pubs PATH] [--procs PATH]

Sem args, usa os CSVs anexados pelo Leonardo na sessão (em %TEMP%).
"""
from __future__ import annotations

import argparse
import csv
import os
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.worksheet import Worksheet


# ============================================================================
# Constants
# ============================================================================

REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_TMP = Path(os.path.expandvars("%TEMP%"))
DEFAULT_PUBS = DEFAULT_TMP / "📬 Publicações 6ee4f13a9ea34506824656a261d99dce_all (4).csv"
DEFAULT_PROCS = DEFAULT_TMP / "⚖️ Processos 78f11cd71d0d4a1e9df58be085e5bb33_all (12).csv"
OUT_XLSX = REPO_ROOT / "logs" / "analise_alertas_round_7.xlsx"

ALERTAS_GRANDES_THRESHOLD = 50  # alertas com >N disparos ganham aba dedicada

# Excel proíbe \\ / * ? : [ ] e limita a 31 chars no nome da aba.
# Mapeia alertas que precisam de nome encurtado / saneado.
SHEET_TITLE_MAP = {
    "Processo/recurso distribuído": "Processo-recurso distribuído",
    "Capturar numeração STJ/TST": "Capturar numeração STJ-TST",
    "Banco do Brasil ausente em partes adversas": "BB ausente em partes adversas",
    "Instância desatualizada (subida)": "Instância desatualizada subida",
}


def _safe_sheet_name(alerta: str) -> str:
    nome = SHEET_TITLE_MAP.get(alerta, alerta)
    for c in "\\/*?:[]":
        nome = nome.replace(c, "-")
    return nome[:31]

# Ordem canônica das 11 abas dedicadas (ordem decrescente confirmada empiricamente)
ABAS_DEDICADAS = [
    "Capturar link externo",
    "Vara desatualizada",
    "Processo não cadastrado",
    "Fase desatualizada (cognitiva)",
    "Processo/recurso distribuído",
    "Turma desatualizada",
    "Banco do Brasil ausente em partes adversas",
    "Incluir julgamento no controle",
    "Capturar numeração STJ/TST",
    "Trânsito em julgado pendente",
    "Instância desatualizada (subida)",
]

RE_CNJ = re.compile(r"^(\d{7}-\d{2}\.\d{4}\.\d\.\d{2}\.\d{4})")
RE_PAGE_ID = re.compile(r"notion\.so/[\w-]+-([0-9a-f]{32})")
RE_VARA_NUMERO = re.compile(r"^(\d+)\s*[ªº]")
RE_TURMA_NUMERO = re.compile(r"(\d+)\s*[ªº]\s*Turma", re.IGNORECASE)


# ============================================================================
# Styling
# ============================================================================

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)

LINK_FONT = Font(name="Calibri", size=11, color="0563C1", underline="single")

# Diagnóstico fills
FILL_VAZIO = PatternFill("solid", fgColor="FCE4E4")  # vermelho claro
FILL_DIVERG = PatternFill("solid", fgColor="FFF2CC")  # amarelo
FILL_OK = PatternFill("solid", fgColor="F2F2F2")  # cinza claro
FILL_REGRA = PatternFill("solid", fgColor="E1D5E7")  # roxo claro (regra suspeita)


def fill_for_diag(label: str) -> PatternFill | None:
    """Retorna fill conforme prefixo do diagnóstico."""
    label_l = label.lower()
    if label_l.startswith("vazio") or label_l.startswith("cadastro vazio"):
        return FILL_VAZIO
    if label_l.startswith("divergente"):
        return FILL_DIVERG
    if label_l.startswith("regra suspeita"):
        return FILL_REGRA
    if label_l.startswith("esperado") or label_l.startswith("ok"):
        return FILL_OK
    return None


# ============================================================================
# Helpers de extração
# ============================================================================

def extrair_cnj(s: str) -> str | None:
    m = RE_CNJ.match(s.strip())
    return m.group(1) if m else None


def extrair_page_id(s: str) -> str | None:
    """Extrai page_id (com hífens) da URL no campo Processo."""
    m = RE_PAGE_ID.search(s)
    if not m:
        return None
    raw = m.group(1)
    return f"{raw[:8]}-{raw[8:12]}-{raw[12:16]}-{raw[16:20]}-{raw[20:32]}"


def extrair_numero_vara(orgao: str) -> str | None:
    """Extrai '9' de '9ª Vara Cível de Brasília' ou '17ª Vara do Trabalho de Brasília - DF'."""
    m = RE_VARA_NUMERO.match(orgao.strip())
    return m.group(1) if m else None


def extrair_numero_turma(orgao: str) -> str | None:
    """Extrai '4' de '4ª Turma Cível' ou '5ª Turma'."""
    m = RE_TURMA_NUMERO.search(orgao.strip())
    return m.group(1) if m else None


def split_multi_select(s: str) -> list[str]:
    return [x.strip() for x in s.split(",") if x.strip()]


# ============================================================================
# Loaders
# ============================================================================

@dataclass
class Pub:
    identif: str
    tribunal: str
    orgao: str
    tipo_com: str
    tipo_doc: str
    classe: str
    cnj: str | None
    proc_page_id: str | None
    cliente_url: str
    partes: str
    texto: str
    data_disp: str
    alertas: list[str] = field(default_factory=list)
    tarefas: list[str] = field(default_factory=list)


@dataclass
class Proc:
    cnj: str
    tribunal: str
    cidade: str
    vara: str
    instancia: str
    fase: str
    natureza: str
    tipo: str
    posicao: str
    turma_2: str
    turma_stj: str
    turma_stf: str
    relator_2: str
    relator_stj: str
    relator_stf: str
    partes_adversas: str
    link_externo: str
    numero_stj: str
    numero_stf: str
    data_transito_cog: str
    status: str


def load_pubs(path: Path) -> list[Pub]:
    pubs: list[Pub] = []
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            cnj = extrair_cnj(row["Processo"])
            page_id = extrair_page_id(row["Processo"])
            pubs.append(
                Pub(
                    identif=row["Identificação"].strip(),
                    tribunal=row["Tribunal"].strip(),
                    orgao=row["Órgão"].strip(),
                    tipo_com=row["Tipo de comunicação"].strip(),
                    tipo_doc=row["Tipo de documento"].strip(),
                    classe=row["Classe"].strip(),
                    cnj=cnj,
                    proc_page_id=page_id,
                    cliente_url=row["Cliente"].strip(),
                    partes=row["Partes"].strip(),
                    texto=row["Texto"].strip(),
                    data_disp=row["Data de disponibilização"].strip(),
                    alertas=split_multi_select(row["Alerta contadoria (app)"]),
                    tarefas=split_multi_select(row["Tarefa sugerida (app)"]),
                )
            )
    return pubs


def load_procs(path: Path) -> dict[str, Proc]:
    procs: dict[str, Proc] = {}
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            cnj = row["Número do processo"].strip()
            if not cnj or cnj.startswith("🧱"):
                continue
            procs[cnj] = Proc(
                cnj=cnj,
                tribunal=row["Tribunal"].strip(),
                cidade=row["Cidade"].strip(),
                vara=row["Vara"].strip(),
                instancia=row["Instância"].strip(),
                fase=row["Fase"].strip(),
                natureza=row["Natureza"].strip(),
                tipo=row["Tipo de processo"].strip(),
                posicao=row["Posição do cliente"].strip(),
                turma_2=row["Turma no 2º grau"].strip(),
                turma_stj=row["Turma no STJ/TST"].strip(),
                turma_stf=row["Turma no STF"].strip(),
                relator_2=row["Relator no 2º grau"].strip(),
                relator_stj=row["Relator no STJ/TST"].strip(),
                relator_stf=row["Relator no STF"].strip(),
                partes_adversas=row["Partes adversas"].strip(),
                link_externo=row["Link externo"].strip(),
                numero_stj=row["Número STJ/TST"].strip(),
                numero_stf=row["Número STF"].strip(),
                data_transito_cog=row["Data do trânsito em julgado (cognitiva)"].strip(),
                status=row["Status"].strip(),
            )
    return procs


# ============================================================================
# Cliente extraction
# ============================================================================

# Nomes adversos típicos para excluir de Pub.Partes
PARTES_ADVERSAS_KNOWN = (
    "BANCO DO BRASIL",
    "BANCO BRADESCO",
    "BRADESCO SAÚDE",
    "BRADESCO SAUDE",
    "PREVI",
    "CASSI",
    "CAIXA DE PREVIDENCIA DOS FUNCIONARIOS",
    "BB ADM",
    "BB ADMINISTRADORA",
    "CAIXA ECONOMICA FEDERAL",
    "CAIXA ECONÔMICA FEDERAL",
)


def extrair_cliente_de_partes(partes: str) -> str:
    """Extrai o(s) nome(s) que NÃO são adversas conhecidas."""
    if not partes:
        return ""
    # Partes vem em formato '[POLO] Nome 1, [POLO] Nome 2, ...'
    pieces = re.split(r",\s*(?=\[)", partes)
    clientes = []
    for p in pieces:
        nome_match = re.match(r"\s*(?:\[[^\]]+\]\s*)?(.+)", p.strip())
        if not nome_match:
            continue
        nome = nome_match.group(1).strip()
        nome_upper = nome.upper()
        if any(adv in nome_upper for adv in PARTES_ADVERSAS_KNOWN):
            continue
        # Ignora pessoas jurídicas óbvias se múltiplas
        clientes.append(nome)
    return "; ".join(clientes[:2])  # limita a 2 nomes


# ============================================================================
# Diagnósticos (um por alerta grande)
# ============================================================================

def diag_capturar_link_externo(pub: Pub, proc: Proc | None) -> tuple[str, str]:
    if proc is None:
        return ("regra suspeita", "Pub sem proc cadastrado — disparo redundante com 'Processo não cadastrado'")
    if not proc.link_externo:
        return ("vazio", "Proc.Link externo em branco — propriedade quase nunca preenchida (1/1107)")
    return ("ok", f"Link cadastrado: {proc.link_externo[:60]}")


def diag_vara_desatualizada(pub: Pub, proc: Proc | None) -> tuple[str, str]:
    if proc is None:
        return ("regra suspeita", "Sem proc cadastrado — alerta deveria ser suprimido")
    if not proc.vara:
        return ("vazio", "Proc.Vara em branco no cadastro — preencher")
    pub_num = extrair_numero_vara(pub.orgao)
    if pub_num and pub_num == proc.vara:
        return ("regra suspeita", f"Cadastro={proc.vara!r} casa com Pub.Órgão={pub.orgao!r} — regra compara strings literais")
    return ("divergente", f"Notion={proc.vara!r} vs Pub.Órgão={pub.orgao!r}")


def diag_processo_nao_cadastrado(pub: Pub, proc: Proc | None) -> tuple[str, str]:
    if proc is not None:
        return ("regra suspeita", f"Proc cadastrado mas alerta disparou — verificar matching CNJ")
    cnj = pub.cnj or "(CNJ ilegível)"
    return ("cadastro pendente", f"CNJ {cnj} não está em ⚖️ Processos — cadastrar")


def diag_fase_cognitiva(pub: Pub, proc: Proc | None) -> tuple[str, str]:
    if proc is None:
        return ("regra suspeita", "Sem proc cadastrado")
    classe = pub.classe.upper()
    if "TRABALHISTA" in classe and proc.fase == "Executiva":
        return ("regra suspeita", "Falso positivo: classe trabalhista permanece após cognitiva — fase já cadastrada como Executiva")
    if "PROCEDIMENTO COMUM" in classe and proc.fase in ("Executiva", "Liquidação de sentença"):
        return ("regra suspeita", f"Falso positivo: classe original do rito vs fase atual={proc.fase}")
    return ("divergente", f"Pub.Classe={pub.classe!r} / Proc.Fase={proc.fase!r}")


def diag_processo_distribuido(pub: Pub, proc: Proc | None) -> tuple[str, str]:
    return ("esperado", "Camada base — Distribuição (Lista qualquer | Intimação Distribuição)")


def diag_turma_desatualizada(pub: Pub, proc: Proc | None) -> tuple[str, str]:
    if proc is None:
        return ("regra suspeita", "Sem proc cadastrado")
    inst = proc.instancia
    if inst == "2º grau":
        cad = proc.turma_2
    elif inst in ("STJ", "TST"):
        cad = proc.turma_stj
    elif inst == "STF":
        cad = proc.turma_stf
    else:
        cad = ""
    if not cad:
        return ("vazio", f"Turma vazia (instância={inst!r}) — preencher")
    pub_num = extrair_numero_turma(pub.orgao)
    cad_num_match = re.match(r"^(\d+)", cad)
    cad_num = cad_num_match.group(1) if cad_num_match else None
    if pub_num and cad_num and pub_num == cad_num:
        return ("regra suspeita", f"Cadastro={cad!r} casa com Pub.Órgão={pub.orgao!r} — regra compara strings literais")
    return ("divergente", f"Inst={inst} | Notion={cad!r} vs Pub.Órgão={pub.orgao!r}")


def diag_bb_ausente(pub: Pub, proc: Proc | None) -> tuple[str, str]:
    if proc is None:
        return ("regra suspeita", "Sem proc cadastrado")
    if "banco do brasil" not in proc.partes_adversas.lower() and "bb" not in proc.partes_adversas.lower():
        return ("vazio", f"Partes adversas={proc.partes_adversas!r} — adicionar Banco do Brasil")
    return ("ok", "BB já cadastrado — verificar texto da pub")


def diag_incluir_julgamento(pub: Pub, proc: Proc | None) -> tuple[str, str]:
    return ("esperado", "Camada base — Edital ou Intimação + Pauta de Julgamento")


def diag_numeracao_stj(pub: Pub, proc: Proc | None) -> tuple[str, str]:
    if proc is None:
        return ("regra suspeita", "Sem proc cadastrado")
    if not proc.numero_stj:
        return ("vazio", f"Pub é do {pub.tribunal} mas Proc.Número STJ/TST vazio")
    return ("ok", f"Número STJ/TST cadastrado: {proc.numero_stj}")


def diag_transito_pendente(pub: Pub, proc: Proc | None) -> tuple[str, str]:
    if proc is None:
        return ("regra suspeita", "Sem proc cadastrado")
    if not proc.data_transito_cog:
        return ("vazio", f"Fase={proc.fase!r}, Data trânsito (cognitiva) vazia — extrair do PJe/memória")
    return ("ok", f"Trânsito={proc.data_transito_cog} cadastrado — investigar contradição")


def diag_subida_instancia(pub: Pub, proc: Proc | None) -> tuple[str, str]:
    if proc is None:
        return ("regra suspeita", "Sem proc cadastrado")
    return ("divergente", f"Pub.Tribunal={pub.tribunal} vs Proc.Instância={proc.instancia!r}")


DIAG_FNS = {
    "Capturar link externo": diag_capturar_link_externo,
    "Vara desatualizada": diag_vara_desatualizada,
    "Processo não cadastrado": diag_processo_nao_cadastrado,
    "Fase desatualizada (cognitiva)": diag_fase_cognitiva,
    "Processo/recurso distribuído": diag_processo_distribuido,
    "Turma desatualizada": diag_turma_desatualizada,
    "Banco do Brasil ausente em partes adversas": diag_bb_ausente,
    "Incluir julgamento no controle": diag_incluir_julgamento,
    "Capturar numeração STJ/TST": diag_numeracao_stj,
    "Trânsito em julgado pendente": diag_transito_pendente,
    "Instância desatualizada (subida)": diag_subida_instancia,
}


# ============================================================================
# Colunas por aba (cada alerta destaca campos relevantes)
# ============================================================================

# (header, getter)  — getter recebe (pub, proc, diag, link, n_red)
COLUNAS_BASE_TODAS = [
    ("CNJ", lambda p, pr, d, lk, nr: p.cnj or ""),
    ("Cliente (de Pub.Partes)", lambda p, pr, d, lk, nr: extrair_cliente_de_partes(p.partes)),
    ("Pub.Tribunal", lambda p, pr, d, lk, nr: p.tribunal),
    ("Pub.Órgão", lambda p, pr, d, lk, nr: p.orgao),
    ("Pub.Tipo com.", lambda p, pr, d, lk, nr: p.tipo_com),
    ("Pub.Tipo doc.", lambda p, pr, d, lk, nr: p.tipo_doc),
    ("Pub.Classe", lambda p, pr, d, lk, nr: p.classe),
    ("Pub.Data disp.", lambda p, pr, d, lk, nr: p.data_disp),
]


def _proc_field(getter):
    def f(p, pr, d, lk, nr):
        return getattr(pr, getter) if pr else "(sem cadastro)"
    return f


COLUNAS_POR_ALERTA: dict[str, list[tuple]] = {
    "Capturar link externo": COLUNAS_BASE_TODAS + [
        ("Proc.Link externo", _proc_field("link_externo")),
        ("Proc.Tribunal", _proc_field("tribunal")),
        ("Proc.Instância", _proc_field("instancia")),
    ],
    "Vara desatualizada": COLUNAS_BASE_TODAS + [
        ("Proc.Vara", _proc_field("vara")),
        ("Proc.Tribunal", _proc_field("tribunal")),
        ("Proc.Cidade", _proc_field("cidade")),
        ("Proc.Instância", _proc_field("instancia")),
    ],
    "Processo não cadastrado": [
        ("CNJ (texto)", lambda p, pr, d, lk, nr: p.cnj or "(CNJ ilegível)"),
        ("Cliente (de Pub.Partes)", lambda p, pr, d, lk, nr: extrair_cliente_de_partes(p.partes)),
        ("Pub.Tribunal", lambda p, pr, d, lk, nr: p.tribunal),
        ("Pub.Órgão", lambda p, pr, d, lk, nr: p.orgao),
        ("Pub.Tipo com.", lambda p, pr, d, lk, nr: p.tipo_com),
        ("Pub.Tipo doc.", lambda p, pr, d, lk, nr: p.tipo_doc),
        ("Pub.Classe", lambda p, pr, d, lk, nr: p.classe),
        ("Pub.Partes (raw)", lambda p, pr, d, lk, nr: p.partes[:120]),
        ("Pub.Texto[:200]", lambda p, pr, d, lk, nr: p.texto[:200]),
    ],
    "Fase desatualizada (cognitiva)": COLUNAS_BASE_TODAS + [
        ("Proc.Fase", _proc_field("fase")),
        ("Proc.Tipo", _proc_field("tipo")),
        ("Proc.Instância", _proc_field("instancia")),
        ("Proc.Natureza", _proc_field("natureza")),
    ],
    "Processo/recurso distribuído": COLUNAS_BASE_TODAS + [
        ("Proc.Instância", _proc_field("instancia")),
        ("Proc.Status", _proc_field("status")),
    ],
    "Turma desatualizada": COLUNAS_BASE_TODAS + [
        ("Proc.Turma 2º", _proc_field("turma_2")),
        ("Proc.Turma STJ/TST", _proc_field("turma_stj")),
        ("Proc.Turma STF", _proc_field("turma_stf")),
        ("Proc.Tribunal", _proc_field("tribunal")),
        ("Proc.Instância", _proc_field("instancia")),
    ],
    "Banco do Brasil ausente em partes adversas": COLUNAS_BASE_TODAS + [
        ("Proc.Partes adversas", _proc_field("partes_adversas")),
        ("Pub.Partes (raw)", lambda p, pr, d, lk, nr: p.partes[:200]),
    ],
    "Incluir julgamento no controle": COLUNAS_BASE_TODAS + [
        ("Proc.Instância", _proc_field("instancia")),
        ("Proc.Fase", _proc_field("fase")),
    ],
    "Capturar numeração STJ/TST": COLUNAS_BASE_TODAS + [
        ("Proc.Número STJ/TST", _proc_field("numero_stj")),
        ("Proc.Tribunal", _proc_field("tribunal")),
        ("Proc.Instância", _proc_field("instancia")),
    ],
    "Trânsito em julgado pendente": COLUNAS_BASE_TODAS + [
        ("Proc.Data trânsito (cog.)", _proc_field("data_transito_cog")),
        ("Proc.Fase", _proc_field("fase")),
    ],
    "Instância desatualizada (subida)": COLUNAS_BASE_TODAS + [
        ("Proc.Instância", _proc_field("instancia")),
        ("Proc.Tribunal", _proc_field("tribunal")),
    ],
}

# Colunas comuns ao final
TAIL_COLS = [
    ("Diagnóstico", None),  # preenchido custom
    ("Detalhe", None),
    ("Pubs detectoras (proc)", None),
    ("Link Notion", None),
]


# ============================================================================
# Geração do xlsx
# ============================================================================

# Larguras mínimas de coluna por header conhecido
COL_WIDTHS = {
    "CNJ": 24,
    "CNJ (texto)": 24,
    "Cliente (de Pub.Partes)": 32,
    "Pub.Tribunal": 10,
    "Pub.Órgão": 38,
    "Pub.Tipo com.": 14,
    "Pub.Tipo doc.": 14,
    "Pub.Classe": 38,
    "Pub.Data disp.": 14,
    "Pub.Partes (raw)": 50,
    "Pub.Texto[:200]": 50,
    "Proc.Link externo": 38,
    "Proc.Tribunal": 10,
    "Proc.Cidade": 14,
    "Proc.Vara": 10,
    "Proc.Instância": 12,
    "Proc.Fase": 18,
    "Proc.Tipo": 14,
    "Proc.Natureza": 12,
    "Proc.Turma 2º": 14,
    "Proc.Turma STJ/TST": 16,
    "Proc.Turma STF": 14,
    "Proc.Partes adversas": 38,
    "Proc.Número STJ/TST": 16,
    "Proc.Data trânsito (cog.)": 18,
    "Proc.Status": 12,
    "Diagnóstico": 18,
    "Detalhe": 50,
    "Pubs detectoras (proc)": 8,
    "Link Notion": 36,
    "Alerta": 36,  # só na aba Outros
}


def _set_widths(ws: Worksheet, headers: list[str]) -> None:
    for idx, h in enumerate(headers, start=1):
        w = COL_WIDTHS.get(h, 16)
        ws.column_dimensions[get_column_letter(idx)].width = w


def _write_header(ws: Worksheet, headers: list[str], start_row: int = 1) -> None:
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def _write_data_row(
    ws: Worksheet,
    row_idx: int,
    pub: Pub,
    proc: Proc | None,
    diag: tuple[str, str],
    link: str,
    n_redundantes: int,
    cols_dyn: list[tuple],
    alerta_label: str | None = None,
) -> None:
    col = 1
    diag_label, diag_detalhe = diag

    # alerta column (apenas na aba Outros)
    if alerta_label is not None:
        ws.cell(row=row_idx, column=col, value=alerta_label)
        col += 1

    for header, getter in cols_dyn:
        v = getter(pub, proc, diag, link, n_redundantes)
        ws.cell(row=row_idx, column=col, value=v)
        col += 1

    # Diagnóstico
    diag_cell = ws.cell(row=row_idx, column=col, value=diag_label)
    fill = fill_for_diag(diag_label)
    if fill:
        diag_cell.fill = fill
    col += 1

    # Detalhe
    ws.cell(row=row_idx, column=col, value=diag_detalhe)
    col += 1

    # Pubs detectoras
    ws.cell(row=row_idx, column=col, value=n_redundantes)
    col += 1

    # Link Notion (hyperlink)
    if link:
        link_cell = ws.cell(row=row_idx, column=col, value=link)
        link_cell.hyperlink = Hyperlink(ref=link_cell.coordinate, target=link)
        link_cell.font = LINK_FONT
    col += 1


def gerar_aba_alerta(
    wb: Workbook,
    titulo: str,
    pubs_filtradas: list[Pub],
    procs: dict[str, Proc],
    diag_fn,
    cols_dyn: list[tuple],
    pubs_por_proc_alerta: dict[tuple[str, str], int],
    alerta_label: str,
) -> int:
    """Cria uma aba e popula com pubs que dispararam um alerta."""
    sheet_name = _safe_sheet_name(titulo)
    ws = wb.create_sheet(title=sheet_name)
    headers = [h for h, _ in cols_dyn] + ["Diagnóstico", "Detalhe", "Pubs detectoras (proc)", "Link Notion"]
    _write_header(ws, headers)
    _set_widths(ws, headers)

    # Ordenar por (Tribunal, Órgão, CNJ)
    pubs_ord = sorted(pubs_filtradas, key=lambda p: (p.tribunal, p.orgao, p.cnj or "", p.identif))

    for i, pub in enumerate(pubs_ord, start=2):
        proc = procs.get(pub.cnj) if pub.cnj else None
        diag = diag_fn(pub, proc)
        n_red = pubs_por_proc_alerta.get((pub.cnj or "", alerta_label), 1)
        link = f"https://www.notion.so/{pub.proc_page_id}" if pub.proc_page_id else ""
        _write_data_row(ws, i, pub, proc, diag, link, n_red, cols_dyn)

    # auto filter
    ws.auto_filter.ref = ws.dimensions
    return len(pubs_ord)


def gerar_aba_outros(
    wb: Workbook,
    pubs: list[Pub],
    procs: dict[str, Proc],
    alertas_pequenos: list[str],
    pubs_por_proc_alerta: dict[tuple[str, str], int],
) -> int:
    ws = wb.create_sheet(title="Outros")
    cols_dyn = [("Alerta", lambda p, pr, d, lk, nr, lab=None: lab)] + COLUNAS_BASE_TODAS + [
        ("Proc.Instância", _proc_field("instancia")),
        ("Proc.Fase", _proc_field("fase")),
        ("Proc.Vara", _proc_field("vara")),
    ]
    headers = [h for h, _ in cols_dyn] + ["Diagnóstico", "Detalhe", "Pubs detectoras (proc)", "Link Notion"]
    _write_header(ws, headers)
    _set_widths(ws, headers)

    rows: list[tuple[Pub, str]] = []
    for pub in pubs:
        for a in pub.alertas:
            if a in alertas_pequenos:
                rows.append((pub, a))
    rows.sort(key=lambda x: (x[1], x[0].tribunal, x[0].orgao, x[0].cnj or ""))

    DIAG_GENERICO = lambda p, pr: ("ok", "Alerta de baixa frequência — ver aba dedicada se houver")

    for i, (pub, alerta) in enumerate(rows, start=2):
        proc = procs.get(pub.cnj) if pub.cnj else None
        # Diagnóstico genérico (alertas pequenos não têm fn dedicada)
        if alerta in DIAG_FNS:
            diag = DIAG_FNS[alerta](pub, proc)
        else:
            diag = _diag_generico(alerta, pub, proc)
        n_red = pubs_por_proc_alerta.get((pub.cnj or "", alerta), 1)
        link = f"https://www.notion.so/{pub.proc_page_id}" if pub.proc_page_id else ""

        col = 1
        ws.cell(row=i, column=col, value=alerta); col += 1
        for header, getter in COLUNAS_BASE_TODAS:
            ws.cell(row=i, column=col, value=getter(pub, proc, diag, link, n_red))
            col += 1
        ws.cell(row=i, column=col, value=proc.instancia if proc else "(sem cadastro)"); col += 1
        ws.cell(row=i, column=col, value=proc.fase if proc else "(sem cadastro)"); col += 1
        ws.cell(row=i, column=col, value=proc.vara if proc else "(sem cadastro)"); col += 1
        diag_cell = ws.cell(row=i, column=col, value=diag[0])
        fill = fill_for_diag(diag[0])
        if fill: diag_cell.fill = fill
        col += 1
        ws.cell(row=i, column=col, value=diag[1]); col += 1
        ws.cell(row=i, column=col, value=n_red); col += 1
        if link:
            lc = ws.cell(row=i, column=col, value=link)
            lc.hyperlink = Hyperlink(ref=lc.coordinate, target=link)
            lc.font = LINK_FONT
        col += 1

    ws.auto_filter.ref = ws.dimensions
    return len(rows)


def _diag_generico(alerta: str, pub: Pub, proc: Proc | None) -> tuple[str, str]:
    """Diagnóstico para alertas pequenos sem função dedicada."""
    if proc is None:
        return ("regra suspeita", "Sem proc cadastrado")
    # Heurísticas simples:
    if "Cidade" in alerta:
        return ("vazio" if not proc.cidade else "divergente", f"Proc.Cidade={proc.cidade!r} | Pub.Órgão={pub.orgao!r}")
    if "Relator" in alerta:
        rel = proc.relator_2 or proc.relator_stj or proc.relator_stf
        return ("vazio" if not rel else "divergente", f"Relator cadastrado={rel!r}")
    if "PREVI" in alerta or "CASSI" in alerta or "BB Adm" in alerta:
        return ("vazio", f"Partes adversas={proc.partes_adversas!r}")
    if "Tribunal fora" in alerta:
        return ("divergente", f"Pub.Tribunal={pub.tribunal!r} fora do vocabulário")
    if "Tema 955" in alerta:
        return ("ok", "Sobrestamento por Tema 955 — verificar manualmente")
    if "natureza" in alerta.lower():
        return ("divergente", f"Pub.Tribunal={pub.tribunal} / Proc.Natureza={proc.natureza}")
    if "Texto imprestável" in alerta:
        return ("ok", "Qualidade do conteúdo DJEN")
    if "arquivado" in alerta.lower():
        return ("divergente", f"Proc.Status={proc.status} mas atividade observada")
    if "Acórdão em processo" in alerta:
        return ("divergente", f"Pub.Tipo doc=Acórdão / Proc.Instância={proc.instancia}")
    if "fase pós-cognitiva" in alerta.lower():
        return ("divergente", f"Pub.Tipo doc={pub.tipo_doc} / Proc.Fase={proc.fase}")
    if "descida" in alerta.lower():
        return ("divergente", f"Pub.Tribunal={pub.tribunal} / Proc.Instância={proc.instancia}")
    if "Vincular cliente" in alerta:
        return ("vazio", "Sinal de cliente do escritório nas Partes mas sem relation")
    if "vinculação" in alerta.lower():
        return ("divergente", "Cliente vinculado ao processo possivelmente errado")
    if "posição" in alerta.lower():
        return ("divergente", f"Proc.Posição={proc.posicao!r}")
    return ("ok", "")


def gerar_resumo(
    wb: Workbook,
    contagens: Counter,
    cobertura: dict[str, Any],
) -> None:
    ws = wb.create_sheet(title="Resumo", index=0)

    # Cabeçalho geral
    ws["A1"] = "Análise dos disparos de alertas — Round 7"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True)
    ws["A2"] = f"Universo: {cobertura['n_pubs']} publicações × {cobertura['n_procs']} processos × {cobertura['n_alertas_distintos']} alertas distintos"
    ws["A3"] = f"Total de disparos: {cobertura['n_disparos']} | Pubs com 2+ alertas: {cobertura['n_2plus']} ({100*cobertura['n_2plus']/cobertura['n_pubs']:.1f}%)"

    # Tabela de alertas
    row = 5
    ws.cell(row=row, column=1, value="Alerta").fill = HEADER_FILL
    ws.cell(row=row, column=1).font = HEADER_FONT
    ws.cell(row=row, column=2, value="Disparos").fill = HEADER_FILL
    ws.cell(row=row, column=2).font = HEADER_FONT
    ws.cell(row=row, column=3, value="% disparos").fill = HEADER_FILL
    ws.cell(row=row, column=3).font = HEADER_FONT
    ws.cell(row=row, column=4, value="Aba").fill = HEADER_FILL
    ws.cell(row=row, column=4).font = HEADER_FONT

    total = sum(contagens.values())
    row += 1
    for alerta, n in contagens.most_common():
        ws.cell(row=row, column=1, value=alerta)
        ws.cell(row=row, column=2, value=n)
        ws.cell(row=row, column=3, value=f"{100*n/total:.1f}%")
        sheet = _safe_sheet_name(alerta) if n > ALERTAS_GRANDES_THRESHOLD else "Outros"
        # link interno
        link_cell = ws.cell(row=row, column=4, value=sheet)
        link_cell.hyperlink = Hyperlink(ref=link_cell.coordinate, target=f"#'{sheet}'!A1", location=f"'{sheet}'!A1")
        link_cell.font = LINK_FONT
        row += 1

    # Legenda de cores
    row += 2
    ws.cell(row=row, column=1, value="Legenda — cores de Diagnóstico").font = Font(bold=True)
    row += 1
    legend = [
        ("vazio / cadastro pendente", FILL_VAZIO, "campo do cadastro está em branco"),
        ("divergente", FILL_DIVERG, "campo populado mas difere da publicação"),
        ("regra suspeita", FILL_REGRA, "candidata a alteração no código"),
        ("esperado / ok", FILL_OK, "alerta correto / sem ação imediata"),
    ]
    for label, fill, desc in legend:
        c = ws.cell(row=row, column=1, value=label)
        c.fill = fill
        ws.cell(row=row, column=2, value=desc)
        row += 1

    # Larguras
    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 36


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pubs", default=str(DEFAULT_PUBS))
    parser.add_argument("--procs", default=str(DEFAULT_PROCS))
    parser.add_argument("--out", default=str(OUT_XLSX))
    args = parser.parse_args()

    pubs = load_pubs(Path(args.pubs))
    procs = load_procs(Path(args.procs))
    print(f"Carregadas {len(pubs)} pubs e {len(procs)} processos")

    # Contagens
    contagens: Counter[str] = Counter()
    pubs_por_proc_alerta: defaultdict[tuple[str, str], int] = defaultdict(int)
    n_2plus = 0
    n_zero = 0
    for pub in pubs:
        if len(pub.alertas) >= 2:
            n_2plus += 1
        if not pub.alertas:
            n_zero += 1
        for a in pub.alertas:
            contagens[a] += 1
            pubs_por_proc_alerta[(pub.cnj or "", a)] += 1

    cobertura = {
        "n_pubs": len(pubs),
        "n_procs": len(procs),
        "n_alertas_distintos": len(contagens),
        "n_disparos": sum(contagens.values()),
        "n_2plus": n_2plus,
        "n_zero": n_zero,
    }

    alertas_grandes = [a for a, n in contagens.items() if n > ALERTAS_GRANDES_THRESHOLD]
    alertas_pequenos = [a for a, n in contagens.items() if n <= ALERTAS_GRANDES_THRESHOLD]
    # Ordena grandes pela ordem de ABAS_DEDICADAS (que coincide com top decrescente)
    alertas_grandes_ord = [a for a in ABAS_DEDICADAS if a in alertas_grandes]
    # Adiciona qualquer outro grande não listado (segurança)
    for a in alertas_grandes:
        if a not in alertas_grandes_ord:
            alertas_grandes_ord.append(a)

    print(f"Alertas com >{ALERTAS_GRANDES_THRESHOLD} disparos: {len(alertas_grandes_ord)} (abas dedicadas)")
    print(f"Alertas com ≤{ALERTAS_GRANDES_THRESHOLD} disparos: {len(alertas_pequenos)} (aba 'Outros')")

    wb = Workbook()
    # Remove default
    wb.remove(wb.active)

    # Resumo primeiro
    gerar_resumo(wb, contagens, cobertura)

    # Abas dedicadas
    for alerta in alertas_grandes_ord:
        pubs_filtradas = [p for p in pubs if alerta in p.alertas]
        diag_fn = DIAG_FNS.get(alerta) or (lambda p, pr: _diag_generico(alerta, p, pr))
        cols = COLUNAS_POR_ALERTA.get(alerta, COLUNAS_BASE_TODAS)
        n = gerar_aba_alerta(
            wb=wb,
            titulo=alerta,
            pubs_filtradas=pubs_filtradas,
            procs=procs,
            diag_fn=diag_fn,
            cols_dyn=cols,
            pubs_por_proc_alerta=pubs_por_proc_alerta,
            alerta_label=alerta,
        )
        print(f"  Aba {alerta!r}: {n} linhas")

    # Aba Outros
    n_outros = gerar_aba_outros(wb, pubs, procs, alertas_pequenos, pubs_por_proc_alerta)
    print(f"  Aba 'Outros': {n_outros} linhas ({len(alertas_pequenos)} alertas distintos)")

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"\nSalvo em: {out}")

    # Métricas para o relatório md
    print("\n=== Métricas para o relatório ===")
    print(f"PROCS_CADASTRO_COMPLETO:")
    print(f"  com Link externo: {sum(1 for p in procs.values() if p.link_externo)}/{len(procs)}")
    print(f"  com Número STJ/TST: {sum(1 for p in procs.values() if p.numero_stj)}/{len(procs)}")
    print(f"  com BB em adversas: {sum(1 for p in procs.values() if 'banco do brasil' in p.partes_adversas.lower())}/{len(procs)}")
    print(f"  com Data trânsito (cog): {sum(1 for p in procs.values() if p.data_transito_cog)}/{len(procs)}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
