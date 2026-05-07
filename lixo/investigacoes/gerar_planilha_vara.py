"""Gera planilha de migração em massa da Vara, no formato consumido pela
aba Importar do app Notion RPADV.

Entrada:
    OneDrive/.../Planilhas projeto Notion/2ª fase/vara_notion_x_datajud.xlsx
    (col 1 = CNJ, col 2 = número da Vara descoberto via DataJud)

Saída:
    logs/migracao_vara_<timestamp>.xlsx
    Headers: page_id | Número do processo | Vara
    Linha 2+: dados, page_id resolvido via cache local (cache.db).

CNJs que não estão no cache (não cadastrados como Processos no Notion)
são listados numa aba auxiliar 'CNJs sem page_id' para revisão.

Uso:
    PYTHONPATH=. python scripts/investigacoes/gerar_planilha_vara.py
"""
from __future__ import annotations

import json
import sqlite3
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SRC_PATH = Path(
    r"C:\Users\LeonardoGuedesdaFons\OneDrive - RICARDO PASSOS ADVOCACIA"
    r"\Área de Trabalho\Planilhas projeto Notion\2ª fase"
    r"\vara_notion_x_datajud.xlsx"
)
CACHE_DB = Path(
    r"C:\Users\LeonardoGuedesdaFons\AppData\Roaming\NotionRPADV\cache.db"
)
OUT_DIR = Path("logs")


def _normalizar_cnj(s: Any) -> str:
    """Strip e mantém só o conteúdo da string original (com máscara)."""
    if s is None:
        return ""
    return str(s).strip()


def _normalizar_vara(s: Any) -> str:
    """Vara é texto no Notion (rich_text). Aceita int, float ou str.

    Float vira int (não queremos '17.0'); int vira string.
    """
    if s is None:
        return ""
    if isinstance(s, float) and s.is_integer():
        return str(int(s))
    return str(s).strip()


def _construir_lookup_cnj_pageid(db: Path) -> dict[str, str]:
    """CNJ (com máscara) → page_id, lendo cache.db do app."""
    conn = sqlite3.connect(str(db))
    try:
        cur = conn.cursor()
        cur.execute(
            "SELECT page_id, data_json FROM records WHERE base = 'Processos'",
        )
        out: dict[str, str] = {}
        for page_id, data_json in cur.fetchall():
            try:
                d = json.loads(data_json)
            except (json.JSONDecodeError, TypeError):
                continue
            cnj = (d.get("numero_do_processo") or "").strip()
            if cnj:
                out[cnj] = page_id
        return out
    finally:
        conn.close()


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]

    if not SRC_PATH.exists():
        print(f"ERRO: planilha de origem não encontrada: {SRC_PATH}")
        return 1

    if not CACHE_DB.exists():
        print(f"ERRO: cache não encontrado: {CACHE_DB}")
        print("       Abra o app Notion RPADV uma vez para popular o cache.")
        return 1

    print(f"Lendo origem: {SRC_PATH.name}")
    wb_src = openpyxl.load_workbook(SRC_PATH, read_only=True, data_only=True)
    ws_src = wb_src[wb_src.sheetnames[0]]

    rows_src: list[tuple[str, str]] = []
    for ri, row in enumerate(ws_src.iter_rows(values_only=True), start=1):
        if ri == 1:
            continue  # cabeçalho
        cnj = _normalizar_cnj(row[0] if len(row) > 0 else None)
        vara = _normalizar_vara(row[1] if len(row) > 1 else None)
        if not cnj:
            continue
        rows_src.append((cnj, vara))
    wb_src.close()
    print(f"  {len(rows_src)} linhas de dados na origem")

    print("Construindo lookup CNJ → page_id a partir do cache…")
    lookup = _construir_lookup_cnj_pageid(CACHE_DB)
    print(f"  {len(lookup)} processos no cache")

    matched: list[tuple[str, str, str]] = []   # page_id, cnj, vara
    missing: list[tuple[str, str]] = []        # cnj, vara (sem page_id)

    for cnj, vara in rows_src:
        page_id = lookup.get(cnj)
        if page_id:
            matched.append((page_id, cnj, vara))
        else:
            missing.append((cnj, vara))

    print(f"  matched: {len(matched)}  /  missing: {len(missing)}")

    # Distribuições amostrais p/ sanity
    sem_vara = [m for m in matched if not m[2]]
    print(f"  matched com Vara vazia: {len(sem_vara)}")

    # ------------------------------------------------------------------
    # Geração do XLSX de saída
    # ------------------------------------------------------------------
    wb_out = openpyxl.Workbook()
    ws_main = wb_out.active
    if ws_main is None:
        print("ERRO: não foi possível criar planilha")
        return 1
    ws_main.title = "Processos"

    # Linha 1: cabeçalhos.
    # 'Número do processo' é obrigatório no validador (validar_linha),
    # mesmo em update por page_id — sem ele a planilha cai com erro
    # "Campo obrigatório 'Número do processo' está ausente ou vazio".
    headers = ["page_id", "Número do processo", "Vara"]
    header_font = Font(color="FFFFFF", bold=True, size=10)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    center = Alignment(horizontal="center", vertical="center")

    for col_idx, h in enumerate(headers, start=1):
        cell = ws_main.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    # Larguras
    ws_main.column_dimensions["A"].width = 38   # page_id (UUID + traços)
    ws_main.column_dimensions["B"].width = 28   # CNJ
    ws_main.column_dimensions["C"].width = 8    # Vara

    # Linhas 2+: dados (formato direto, sem linha de dica — o importar.py
    # lê headers da linha 1 e dados da linha 2 em diante)
    for ri, (page_id, cnj, vara) in enumerate(matched, start=2):
        ws_main.cell(row=ri, column=1, value=page_id)
        ws_main.cell(row=ri, column=2, value=cnj)
        ws_main.cell(row=ri, column=3, value=vara)

    # Congela cabeçalho
    ws_main.freeze_panes = "A2"

    # Aba auxiliar: CNJs não encontrados no cache (não migrar — revisar)
    if missing:
        ws_miss = wb_out.create_sheet("CNJs sem page_id")
        ws_miss.cell(row=1, column=1, value="Número do processo").font = header_font
        ws_miss.cell(row=1, column=1).fill = header_fill
        ws_miss.cell(row=1, column=1).alignment = center
        ws_miss.cell(row=1, column=2, value="Vara segundo DataJud").font = header_font
        ws_miss.cell(row=1, column=2).fill = header_fill
        ws_miss.cell(row=1, column=2).alignment = center
        ws_miss.column_dimensions["A"].width = 28
        ws_miss.column_dimensions["B"].width = 22
        for ri, (cnj, vara) in enumerate(missing, start=2):
            ws_miss.cell(row=ri, column=1, value=cnj)
            ws_miss.cell(row=ri, column=2, value=vara)
        ws_miss.freeze_panes = "A2"

    # Aba de leitura: instruções rápidas de uso
    ws_inst = wb_out.create_sheet("Instruções")
    ws_inst.column_dimensions["A"].width = 90
    titulo_font = Font(bold=True, size=13, color="1F4E79")
    normal_font = Font(size=10)
    instrucoes: list[tuple[str, Any]] = [
        ("Migração em massa — Vara (gerada automaticamente)", titulo_font),
        ("", normal_font),
        ("Como usar:", Font(bold=True, size=10)),
        ("1. Abra o app Notion RPADV.", normal_font),
        ("2. Vá em Importar.", normal_font),
        ("3. Base de dados: 'Processos'.", normal_font),
        ("4. Escolher arquivo (.xlsx) → selecione esta planilha.", normal_font),
        ("5. Confira a pré-visualização e clique em Importar.", normal_font),
        ("", normal_font),
        ("Detalhes do formato:", Font(bold=True, size=10)),
        (
            "• Aba 'Processos' — formato consumido pelo app: row 1 = headers, "
            "row 2+ = dados.", normal_font,
        ),
        (
            "• Coluna 'page_id' = ID da página Notion já existente. "
            "Garante atualização (não cria duplicata).",
            normal_font,
        ),
        (
            "• Coluna 'Número do processo' = CNJ do processo. Obrigatória "
            "para o validador da aba Importar (mesmo em update por page_id).",
            normal_font,
        ),
        ("• Coluna 'Vara' = novo valor a aplicar (texto, ex: '17').", normal_font),
        ("", normal_font),
        (
            f"Origem: {SRC_PATH.name}", Font(italic=True, size=9, color="595959"),
        ),
        (
            f"Linhas migráveis: {len(matched)}  |  Linhas sem page_id: {len(missing)}",
            Font(italic=True, size=9, color="595959"),
        ),
        (
            f"Gerada em: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
            Font(italic=True, size=9, color="595959"),
        ),
    ]
    if missing:
        instrucoes.insert(
            -3,
            (
                "⚠ Aba 'CNJs sem page_id': processos da origem que NÃO estão "
                "cadastrados no cache local (provavelmente não estão na base "
                "Processos do Notion). Revise antes de migrar.",
                Font(italic=True, color="9C5700", size=10),
            ),
        )
    for ri, (texto, fnt) in enumerate(instrucoes, start=1):
        cell = ws_inst.cell(row=ri, column=1, value=texto)
        cell.font = fnt
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws_inst.row_dimensions[ri].height = 18 if ri > 1 else 24

    # Salva
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y-%m-%d-%H%M")
    out_path = OUT_DIR / f"migracao_vara_{ts}.xlsx"
    wb_out.save(str(out_path))
    print(f"\nGerada: {out_path}")
    print(f"  Aba 'Processos':       {len(matched)} linha(s) migráveis")
    if missing:
        print(f"  Aba 'CNJs sem page_id': {len(missing)} linha(s) p/ revisão")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
