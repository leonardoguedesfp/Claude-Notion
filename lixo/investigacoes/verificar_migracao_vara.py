"""Verifica se a migração em massa de Vara foi bem-sucedida e isolada.

Compara três fontes:
1. Origem (vara_notion_x_datajud.xlsx) — intenção do usuário (CNJ, Vara desejada)
2. Audit XLSX (datajud_consulta_2026-05-05-1529.xlsx) — snapshot PRÉ-migração
   (timestamp imediatamente antes do import; tem todas as ▸ atual)
3. CSV pós-migração — estado atual exportado do Notion

Verificações:
A. Para cada CNJ na origem: CSV.Vara == origem.vara_desejada (sucesso da migração)
B. Para cada CNJ na origem: CSV[outras_props] == audit[outras_props_atual]
   exceto: Vara (intencional), Atualizado em (esperado bump)
C. Para CNJs NÃO na origem (que não deviam ter mudado): mesma checagem B.

Uso:
    PYTHONPATH=. python scripts/investigacoes/verificar_migracao_vara.py
"""
from __future__ import annotations

import csv
import re
import sys
from pathlib import Path
from typing import Any

import openpyxl


SRC_PATH = Path(
    r"C:\Users\LeonardoGuedesdaFons\OneDrive - RICARDO PASSOS ADVOCACIA"
    r"\Área de Trabalho\Planilhas projeto Notion\2ª fase"
    r"\vara_notion_x_datajud.xlsx"
)
AUDIT_XLSX = Path("logs/datajud_consulta_2026-05-05-1529.xlsx")
CSV_POS = Path(
    r"C:\Users\LeonardoGuedesdaFons\AppData\Local\Temp"
    r"\⚖️ Processos 78f11cd71d0d4a1e9df58be085e5bb33_all (13).csv"
)


# Colunas do audit XLSX (▸ atual) → coluna do CSV pós
AUDIT_TO_CSV: dict[str, str] = {
    "Número do processo ▸ atual":               "Número do processo",
    "Tribunal ▸ atual":                         "Tribunal",
    "Instância ▸ atual":                        "Instância",
    "Vara ▸ atual":                             "Vara",
    "Cidade ▸ atual":                           "Cidade",
    "Data de distribuição ▸ atual":             "Data de distribuição",
    "Data do trânsito em julgado (cognitiva) ▸ atual":
        "Data do trânsito em julgado (cognitiva)",
    "Status ▸ atual":                           "Status",
    "Fase ▸ atual":                             "Fase",
    "Número STJ/TST ▸ atual":                   "Número STJ/TST",
    "Turma no 2º grau ▸ atual":                 "Turma no 2º grau",
    "Turma no STJ/TST ▸ atual":                 "Turma no STJ/TST",
    "Relator no 2º grau ▸ atual":               "Relator no 2º grau",
    "Relator no STJ/TST ▸ atual":               "Relator no STJ/TST",
    "Tema 955 — Sobrestado ▸ atual":            "Tema 955 — Sobrestado",
    "Clientes ▸ atual":                         "Clientes",
    "Detalhamento da ação ▸ atual":             "Detalhamento da ação",
    "Tipo de ação ▸ atual":                     "Tipo de ação",
    "Desdobramentos ▸ atual":                   "Desdobramentos",
    "Tarefas ▸ atual":                          "Tarefas",
    "Natureza ▸ atual":                         "Natureza",
    "Turma no STF ▸ atual":                     "Turma no STF",
    "Observações ▸ atual":                      "Observações",
    "Sobrestado - IRR 20 ▸ atual":              "Sobrestado - IRR 20",
    "ID Legal One ▸ atual":                     "ID Legal One",
    "Link externo ▸ atual":                     "Link externo",
    "Partes adversas ▸ atual":                  "Partes adversas",
    "Data do trânsito em julgado (executiva) ▸ atual":
        "Data do trânsito em julgado (executiva)",
    "Número STF ▸ atual":                       "Número STF",
    "Sobrestado - TJ conexa ▸ atual":           "Sobrestado - TJ conexa",
    "Relator no STF ▸ atual":                   "Relator no STF",
    "Processo pai ▸ atual":                     "Processo pai",
    "Posição do cliente ▸ atual":               "Posição do cliente",
    "Documentos ▸ atual":                       "Documentos",
    "Publicações ▸ atual":                      "Publicações",
    "Tipo de processo ▸ atual":                 "Tipo de processo",
}

# Colunas que esperamos divergir e queremos pular na comparação:
#  • "Vara" — é justamente o objetivo da migração (mudou de propósito)
#  • "Atualizado em" — Notion bumpa sempre que update_page é chamado
#  • CSV "Criado em" / audit não tem coluna direta — pular se não estiver no mapa
COLS_PULAR = {"Vara"}


_MESES_PT = {
    "janeiro": "01", "fevereiro": "02", "março": "03", "abril": "04",
    "maio": "05", "junho": "06", "julho": "07", "agosto": "08",
    "setembro": "09", "outubro": "10", "novembro": "11", "dezembro": "12",
}

# Captura UUIDs (com ou sem hífen) — o CSV exporta relations como
# "Title (https://www.notion.so/Title-<uuid_sem_hifen>?pvs=21)" e o
# audit XLSX exporta como UUIDs separados por vírgula. Normalizamos
# ambos extraindo só os UUIDs em ordem alfabética.
_UUID_RE = re.compile(r"[0-9a-f]{8}-?[0-9a-f]{4}-?[0-9a-f]{4}-?[0-9a-f]{4}-?[0-9a-f]{12}")


def _extrair_uuids(s: str) -> list[str]:
    raw = _UUID_RE.findall(s)
    norm = ["".join(c for c in u.lower() if c != "-") for u in raw]
    return sorted(set(norm))


def _norm_data(s: str) -> str:
    """Aceita 'YYYY-MM-DD', 'D de mês de YYYY', 'D de mês de YYYY HH:MM'.

    Devolve YYYY-MM-DD ou string original se não casou.
    """
    s = s.strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
        return s
    m = re.match(
        r"(\d{1,2}) de (janeiro|fevereiro|março|abril|maio|junho|julho|"
        r"agosto|setembro|outubro|novembro|dezembro) de (\d{4})",
        s, re.IGNORECASE,
    )
    if m:
        d = m.group(1).zfill(2)
        mo = _MESES_PT[m.group(2).lower()]
        y = m.group(3)
        return f"{y}-{mo}-{d}"
    return s


def _norm(v: Any) -> str:
    """Normalização defensiva para comparar valores entre planilhas e CSV.

    1. None → ""; trim
    2. "17.0" → "17" (float-from-int comum em xlsx)
    3. Vírgulas: "A,B" → "A, B"
    """
    if v is None:
        return ""
    s = str(v).strip()
    if re.fullmatch(r"\d+\.0", s):
        s = s[:-2]
    s = re.sub(r"\s*,\s*", ", ", s)
    return s


def _comparar(col: str, pre: str, pos: str) -> bool:
    """True se pre == pos depois de normalização específica da coluna.

    - Datas: compara YYYY-MM-DD após parse
    - Relations (Clientes, Tarefas, Desdobramentos, Processo pai,
      Documentos, Publicações): compara conjuntos de UUIDs
    - Multi_select (Tipo de ação, Detalhamento, Natureza,
      Posição do cliente, Tipo de processo): compara conjuntos de tokens
    - Resto: comparação direta
    """
    if pre == pos:
        return True

    # Datas
    if "Data" in col or "trânsito" in col.lower():
        return _norm_data(pre) == _norm_data(pos)

    # Relations: extrai UUIDs e compara como conjuntos
    rels = {
        "Clientes", "Tarefas", "Desdobramentos", "Processo pai",
        "Documentos", "Publicações",
    }
    if col in rels:
        return _extrair_uuids(pre) == _extrair_uuids(pos)

    # Multi-select: ordem pode variar
    multis = {
        "Tipo de ação", "Detalhamento da ação", "Natureza",
        "Posição do cliente", "Tipo de processo", "Partes adversas",
    }
    if col in multis:
        a = sorted(t.strip() for t in pre.split(",") if t.strip())
        b = sorted(t.strip() for t in pos.split(",") if t.strip())
        return a == b

    return False


def _ler_origem(path: Path) -> dict[str, str]:
    """CNJ → Vara desejada."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    out: dict[str, str] = {}
    for ri, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if ri == 1:
            continue
        cnj = (str(row[0]) if row and row[0] else "").strip()
        vara = row[1] if len(row) > 1 else None
        if cnj:
            out[cnj] = _norm(vara)
    wb.close()
    return out


def _ler_audit(path: Path) -> dict[str, dict[str, str]]:
    """page_id → {col_csv_equivalente: valor_pre_norm}."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["DataJUD"]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(c) if c is not None else "" for c in rows[0]]
    out: dict[str, dict[str, str]] = {}
    for r in rows[1:]:
        d = {headers[i]: r[i] for i in range(len(headers)) if i < len(r)}
        page_id = (d.get("page_id") or "").strip()
        if not page_id:
            continue
        snapshot: dict[str, str] = {}
        for audit_col, csv_col in AUDIT_TO_CSV.items():
            snapshot[csv_col] = _norm(d.get(audit_col))
        out[page_id] = snapshot
    wb.close()
    return out


def _ler_csv_pos(path: Path) -> dict[str, dict[str, str]]:
    """CNJ → {col: valor_pos_norm}."""
    out: dict[str, dict[str, str]] = {}
    with path.open(encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            cnj_key = next(
                (k for k in row.keys() if k.lstrip("﻿").strip() == "Número do processo"),
                None,
            )
            if cnj_key is None:
                continue
            cnj = (row[cnj_key] or "").strip()
            if not cnj:
                continue
            normalized: dict[str, str] = {
                k.lstrip("﻿").strip(): _norm(v) for k, v in row.items()
            }
            out[cnj] = normalized
    return out


def _ler_lookup_cnj_pageid_audit(path: Path) -> dict[str, str]:
    """A partir do audit XLSX, monta CNJ → page_id."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["DataJUD"]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(c) if c is not None else "" for c in rows[0]]
    out: dict[str, str] = {}
    for r in rows[1:]:
        d = {headers[i]: r[i] for i in range(len(headers)) if i < len(r)}
        cnj = (d.get("Número do processo ▸ atual") or "").strip()
        page_id = (d.get("page_id") or "").strip()
        if cnj and page_id:
            out[cnj] = page_id
    wb.close()
    return out


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]

    print("Lendo origem (intenção)…")
    origem = _ler_origem(SRC_PATH)
    print(f"  {len(origem)} CNJs na origem")

    print("Lendo audit XLSX (snapshot pré)…")
    audit_por_pageid = _ler_audit(AUDIT_XLSX)
    cnj_to_pageid = _ler_lookup_cnj_pageid_audit(AUDIT_XLSX)
    print(f"  {len(audit_por_pageid)} processos no snapshot")

    print("Lendo CSV pós-migração…")
    csv_por_cnj = _ler_csv_pos(CSV_POS)
    print(f"  {len(csv_por_cnj)} processos no CSV")

    print()

    # ------------------------------------------------------------------
    # Verificação A — Vara migrada para o valor desejado
    # ------------------------------------------------------------------
    print("=" * 70)
    print("A) Vara migrada para o valor pedido (origem → CSV pós)")
    print("=" * 70)

    vara_ok = 0
    vara_fail: list[tuple[str, str, str]] = []   # cnj, esperado, atual
    vara_orphans: list[str] = []                 # CNJ na origem mas não no CSV

    for cnj, vara_esperada in origem.items():
        csv_row = csv_por_cnj.get(cnj)
        if csv_row is None:
            vara_orphans.append(cnj)
            continue
        vara_atual = csv_row.get("Vara", "")
        if vara_atual == vara_esperada:
            vara_ok += 1
        else:
            vara_fail.append((cnj, vara_esperada, vara_atual))

    print(f"  OK:                {vara_ok}/{len(origem)}")
    print(f"  Vara divergente:   {len(vara_fail)}")
    print(f"  CNJ órfão (no CSV): {len(vara_orphans)}")
    if vara_fail:
        print("  Primeiras 10 divergências:")
        for cnj, esp, atu in vara_fail[:10]:
            print(f"    {cnj}  esperado={esp!r}  atual={atu!r}")
    if vara_orphans:
        print("  Primeiros 5 órfãos:")
        for cnj in vara_orphans[:5]:
            print(f"    {cnj}")

    # ------------------------------------------------------------------
    # Verificação B — outras propriedades NÃO mudaram nos 669 alvos
    # ------------------------------------------------------------------
    print()
    print("=" * 70)
    print("B) Outras propriedades INTACTAS nos 669 alvos da migração")
    print("=" * 70)

    drift: dict[str, list[tuple[str, str, str]]] = {}  # col → [(cnj, pre, pos)]
    n_cmp = 0
    sem_audit: list[str] = []

    for cnj in origem:
        page_id = cnj_to_pageid.get(cnj)
        if page_id is None:
            sem_audit.append(cnj)
            continue
        pre = audit_por_pageid.get(page_id)
        pos = csv_por_cnj.get(cnj)
        if pre is None or pos is None:
            sem_audit.append(cnj)
            continue
        for col in AUDIT_TO_CSV.values():
            if col in COLS_PULAR:
                continue
            v_pre = pre.get(col, "")
            v_pos = pos.get(col, "")
            if not _comparar(col, v_pre, v_pos):
                drift.setdefault(col, []).append((cnj, v_pre, v_pos))
            n_cmp += 1

    print(f"  Comparações realizadas: {n_cmp}")
    print(f"  CNJs sem snapshot pré:  {len(sem_audit)}")
    print(f"  Colunas com drift:      {len(drift)}")
    if drift:
        for col, items in sorted(drift.items(), key=lambda kv: -len(kv[1])):
            print(f"\n  Coluna {col!r}: {len(items)} divergência(s)")
            for cnj, pre, pos in items[:5]:
                print(f"    {cnj}")
                print(f"      pre:  {pre[:120]!r}")
                print(f"      pos:  {pos[:120]!r}")
            if len(items) > 5:
                print(f"    … (+{len(items) - 5} mais)")

    # ------------------------------------------------------------------
    # Verificação C — CNJs FORA da migração também não mudaram
    # ------------------------------------------------------------------
    print()
    print("=" * 70)
    print("C) Processos NÃO migrados também ficaram intactos (incluindo Vara)")
    print("=" * 70)

    fora = [cnj for cnj in csv_por_cnj if cnj not in origem]
    drift_fora: dict[str, list[tuple[str, str, str]]] = {}
    n_cmp_fora = 0
    for cnj in fora:
        page_id = cnj_to_pageid.get(cnj)
        if page_id is None:
            continue
        pre = audit_por_pageid.get(page_id)
        pos = csv_por_cnj.get(cnj)
        if pre is None or pos is None:
            continue
        for col in AUDIT_TO_CSV.values():
            v_pre = pre.get(col, "")
            v_pos = pos.get(col, "")
            if not _comparar(col, v_pre, v_pos):
                drift_fora.setdefault(col, []).append((cnj, v_pre, v_pos))
            n_cmp_fora += 1

    print(f"  Processos fora da migração: {len(fora)}")
    print(f"  Comparações realizadas:     {n_cmp_fora}")
    print(f"  Colunas com drift:          {len(drift_fora)}")
    if drift_fora:
        for col, items in sorted(drift_fora.items(), key=lambda kv: -len(kv[1])):
            print(f"\n  Coluna {col!r}: {len(items)} divergência(s)")
            for cnj, pre, pos in items[:5]:
                print(f"    {cnj}")
                print(f"      pre:  {pre[:120]!r}")
                print(f"      pos:  {pos[:120]!r}")
            if len(items) > 5:
                print(f"    … (+{len(items) - 5} mais)")

    # ------------------------------------------------------------------
    # Veredito
    # ------------------------------------------------------------------
    print()
    print("=" * 70)
    print("VEREDITO")
    print("=" * 70)
    sucesso_vara = vara_ok == len(origem) and not vara_fail and not vara_orphans
    sem_drift = not drift and not drift_fora
    if sucesso_vara and sem_drift:
        print("✓ Migração 100% bem-sucedida e isolada.")
        return 0
    if sucesso_vara and drift:
        print("⚠ Vara foi atualizada corretamente, mas há drift em outras "
              "colunas — investigar lista acima.")
        return 2
    if not sucesso_vara:
        print("✗ Vara NÃO foi totalmente migrada — investigar lista de divergências.")
        return 3
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
