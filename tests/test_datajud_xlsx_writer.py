"""Testes do writer xlsx (Componente 3).

Cobre estrutura de abas, ordem de colunas, **invariante de célula vazia
via inspeção do XML interno do xlsx** (não via openpyxl, que não
distingue célula `""` de célula não-escrita), coloração por confiança,
DataValidation, freeze panes e meta column.
"""
from __future__ import annotations

import json
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from notion_rpadv.services.datajud_enricher import (
    REGRAS_ORIGEM,
    ResultadoEnriquecimento,
)
from notion_rpadv.services.datajud_xlsx_writer import (
    CELL_FILL_DIVERGENTE_OU_BAIXA_HEX,
    CELL_FILL_VAZIO_SUGERIDO_ALTA_HEX,
    HEADER_FILL_HIGH_HEX,
    HEADER_FILL_LOW_HEX,
    SHEET_NAME_DATAJUD,
    SHEET_NAME_INSTRUCOES,
    SHEET_NAME_VOCABULARIOS,
    STATUS_DIVERGENCIA,
    gerar_xlsx,
    status_divergencia,
)


# ---------------------------------------------------------------------------
# Mock PropSpec
# ---------------------------------------------------------------------------


@dataclass
class _SpecMock:
    """Imita ``notion_bulk_edit.schemas.PropSpec`` com os campos que o
    writer realmente lê."""
    notion_name: str
    tipo: str
    label: str = ""
    editavel: bool = True
    obrigatorio: bool = False
    opcoes: tuple[str, ...] = field(default_factory=tuple)


def _spec(
    notion_name: str,
    tipo: str = "rich_text",
    *,
    opcoes: tuple[str, ...] = (),
) -> _SpecMock:
    return _SpecMock(
        notion_name=notion_name,
        tipo=tipo,
        label=notion_name,
        opcoes=opcoes,
    )


def _schema_processos_minimo() -> dict[str, _SpecMock]:
    """Schema mock cobrindo as 14 enriquecidas + 3 auxiliares (suficiente
    para exercitar o writer)."""
    out: dict[str, _SpecMock] = {}
    # 14 enriquecidas — em ordem do REGRAS_ORIGEM, com tipos compatíveis
    out["numero_do_processo"] = _spec("Número do processo", "title")
    out["tribunal"]            = _spec(
        "Tribunal", "select",
        opcoes=("TJDFT", "TRT/10", "TRT/2", "TST", "STJ"),
    )
    out["instancia"]           = _spec(
        "Instância", "select",
        opcoes=("1º grau", "2º grau", "TST", "STJ", "STF"),
    )
    out["vara"]                = _spec("Vara", "rich_text")
    out["cidade"]              = _spec("Cidade", "rich_text")
    out["data_distribuicao"]   = _spec("Data de distribuição", "date")
    out["data_transito_cog"]   = _spec(
        "Data do trânsito em julgado (cognitiva)", "date",
    )
    out["status"]              = _spec(
        "Status", "select",
        opcoes=("Ativo", "Arquivado provisoriamente (tema 955)", "Arquivado"),
    )
    out["fase"]                = _spec(
        "Fase", "select",
        opcoes=("Cognitiva", "Liquidação pendente",
                "Liquidação de sentença", "Executiva",
                "TJ - sentença não será executada"),
    )
    out["numero_stj"]          = _spec("Número STJ", "rich_text")
    out["turma_2g"]            = _spec("Turma no 2º grau", "rich_text")
    out["turma_stj_tst"]       = _spec("Turma no STJ/TST", "rich_text")
    out["relator_2g"]          = _spec("Relator no 2º grau", "rich_text")
    out["relator_stj_tst"]     = _spec("Relator no STJ/TST", "rich_text")
    # 3 auxiliares (não cobertas pelo enricher)
    out["partes_adversas"]     = _spec("Partes adversas", "multi_select")
    out["link_externo"]        = _spec("Link externo", "url")
    out["observacoes"]         = _spec("Observações", "rich_text")
    return out


def _resultado(
    *,
    page_id: str = "abc123",
    cnj: str = "0000123-45.2024.5.10.0001",
    diagnostico: str = "OK",
    propriedades: dict[str, Any] | None = None,
    fontes: list[str] | None = None,
    movs_brutos: dict[str, list[dict[str, Any]]] | None = None,
) -> ResultadoEnriquecimento:
    if propriedades is None:
        propriedades = {r.nome_notion: None for r in REGRAS_ORIGEM}
    return ResultadoEnriquecimento(
        numero_cnj=cnj,
        page_id=page_id,
        diagnostico=diagnostico,
        propriedades_sugeridas=propriedades,
        fontes_tribunal=fontes if fontes is not None else ["trt10"],
        movimentos_brutos_por_grau=movs_brutos or {"G1": [{"codigo": 26}]},
    )


# ---------------------------------------------------------------------------
# 1) Estrutura de abas
# ---------------------------------------------------------------------------


def test_estrutura_de_abas(tmp_path: Path) -> None:
    """3 abas: DataJUD (visível), Instruções (visível), _vocabularios (oculta)."""
    out = tmp_path / "test.xlsx"
    schema = _schema_processos_minimo()
    gerar_xlsx([_resultado()], processos_cache={}, schema=schema, output_path=out)

    wb = load_workbook(out)
    assert SHEET_NAME_DATAJUD in wb.sheetnames
    assert SHEET_NAME_INSTRUCOES in wb.sheetnames
    assert SHEET_NAME_VOCABULARIOS in wb.sheetnames
    assert wb[SHEET_NAME_VOCABULARIOS].sheet_state == "hidden"
    assert wb[SHEET_NAME_DATAJUD].sheet_state == "visible"
    assert wb[SHEET_NAME_INSTRUCOES].sheet_state == "visible"


# ---------------------------------------------------------------------------
# 2) Ordem de colunas
# ---------------------------------------------------------------------------


def test_ordem_de_colunas(tmp_path: Path) -> None:
    """Coluna A=page_id, B=Diagnóstico, C..= 14 enriquecidas × 3 colunas
    (importável, ▸ atual, ▸ status), depois auxiliares ▸ atual,
    e por último __datajud_meta."""
    out = tmp_path / "test.xlsx"
    schema = _schema_processos_minimo()
    gerar_xlsx([_resultado()], processos_cache={}, schema=schema, output_path=out)

    wb = load_workbook(out)
    ws = wb[SHEET_NAME_DATAJUD]

    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]

    # Estrutura esperada
    assert headers[0] == "page_id"
    assert headers[1] == "Diagnóstico"
    # 14 enriquecidas × 3 colunas
    col = 2
    for regra in REGRAS_ORIGEM:
        assert headers[col]     == regra.nome_notion,            f"col {col}"
        assert headers[col + 1] == f"{regra.nome_notion} ▸ atual"
        assert headers[col + 2] == f"{regra.nome_notion} ▸ status"
        col += 3
    # 3 auxiliares (▸ atual)
    auxiliares_esperadas = ["Partes adversas", "Link externo", "Observações"]
    for nn in auxiliares_esperadas:
        assert headers[col] == f"{nn} ▸ atual"
        col += 1
    # __datajud_meta no final
    assert headers[-1] == "__datajud_meta"


# ---------------------------------------------------------------------------
# 3) INVARIANTE CRÍTICA — célula vazia nunca escrita (via XML interno)
# ---------------------------------------------------------------------------


def test_celula_vazia_nao_escrita_invariante(tmp_path: Path) -> None:
    """Inspeciona o XML interno do xlsx (não basta openpyxl).

    Cenário: enriquecimento traz Tribunal e Vara, mas Cidade=None,
    Status=None etc. As colunas dessas propriedades **não devem**
    existir no XML do sheet (nem via shared strings vazia, nem via
    inline string vazia).
    """
    out = tmp_path / "test.xlsx"
    schema = _schema_processos_minimo()

    # Resultado com várias propriedades None — gatilho da invariante
    propriedades = {r.nome_notion: None for r in REGRAS_ORIGEM}
    propriedades["Tribunal"] = "TRT/10"
    propriedades["Vara"] = "13"
    res = _resultado(propriedades=propriedades)

    # Cache também com algumas chaves vazias
    processos_cache: dict[str, dict[str, Any]] = {
        "abc123": {
            "tribunal": "TRT/10",
            "vara": None,             # vazio no Notion
            "cidade": "",             # string vazia no Notion (cuidado!)
            "partes_adversas": None,
        },
    }

    gerar_xlsx([res], processos_cache=processos_cache, schema=schema, output_path=out)

    with zipfile.ZipFile(out) as zf:
        # 1) Verifica todas as sheets de dados
        sheet_paths = [n for n in zf.namelist() if n.startswith("xl/worksheets/")]
        for sp in sheet_paths:
            xml = zf.read(sp).decode("utf-8")
            # <v></v> ou <v/> indicariam célula com valor explicitamente vazio
            assert "<v></v>" not in xml, f"<v></v> em {sp}"
            assert "<v/>" not in xml, f"<v/> em {sp}"
            # Inline string vazia: <is><t></t></is> ou <is><t/></is>
            assert "<is><t></t></is>" not in xml, f"<is><t></t></is> em {sp}"
            assert "<is><t/></is>" not in xml, f"<is><t/></is> em {sp}"

        # 2) Shared strings — qualquer <si><t...></t></si> com texto vazio
        if "xl/sharedStrings.xml" in zf.namelist():
            ss_xml = zf.read("xl/sharedStrings.xml").decode("utf-8")
            assert re.search(r"<t[^>]*></t>", ss_xml) is None, (
                "string compartilhada vazia em sharedStrings.xml"
            )
            assert "<t/>" not in ss_xml, "<t/> em sharedStrings.xml"


# ---------------------------------------------------------------------------
# 4) 24 colunas auxiliares ocultas (3 no schema mínimo de teste)
# ---------------------------------------------------------------------------


def test_colunas_auxiliares_estao_ocultas(tmp_path: Path) -> None:
    """Auxiliares (▸ atual fora do escopo) são todas hidden=True."""
    out = tmp_path / "test.xlsx"
    schema = _schema_processos_minimo()
    gerar_xlsx([_resultado()], processos_cache={}, schema=schema, output_path=out)

    wb = load_workbook(out)
    ws = wb[SHEET_NAME_DATAJUD]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    auxiliares_esperadas = {"Partes adversas", "Link externo", "Observações"}

    for c, h in enumerate(headers, start=1):
        if h is None:
            continue
        nn_aux = h.replace(" ▸ atual", "")
        if nn_aux in auxiliares_esperadas and h.endswith(" ▸ atual"):
            from openpyxl.utils import get_column_letter
            letra = get_column_letter(c)
            assert ws.column_dimensions[letra].hidden is True, (
                f"coluna {letra} ({h}) deveria estar oculta"
            )


# ---------------------------------------------------------------------------
# 5) Coloração — alta confiança (verde claro) em vazio_sugerido
# ---------------------------------------------------------------------------


def test_coloracao_celulas_alta_confianca(tmp_path: Path) -> None:
    """Notion vazio + DataJud sugere Vara='13' (alta confiança) →
    célula da coluna 'Vara' fica com fill verde claro D4E7D4."""
    out = tmp_path / "test.xlsx"
    schema = _schema_processos_minimo()

    propriedades = {r.nome_notion: None for r in REGRAS_ORIGEM}
    propriedades["Vara"] = "13"
    res = _resultado(propriedades=propriedades)

    # cache vazio pra "Vara" → vazio_sugerido (alta)
    processos_cache: dict[str, dict[str, Any]] = {"abc123": {"vara": None}}

    gerar_xlsx([res], processos_cache=processos_cache, schema=schema, output_path=out)

    wb = load_workbook(out)
    ws = wb[SHEET_NAME_DATAJUD]
    # Encontra a coluna importável "Vara"
    col_vara: int | None = None
    for c in range(1, ws.max_column + 1):
        if ws.cell(row=1, column=c).value == "Vara":
            col_vara = c
            break
    assert col_vara is not None
    cell = ws.cell(row=2, column=col_vara)
    assert cell.value == "13"
    fill = cell.fill
    assert fill is not None
    # Verde claro
    assert fill.fgColor.value.upper().endswith(
        CELL_FILL_VAZIO_SUGERIDO_ALTA_HEX,
    ), f"fill cor inesperada: {fill.fgColor.value}"


# ---------------------------------------------------------------------------
# 6) Coloração — amarela em divergente (alta) + vazio_sugerido (baixa)
# ---------------------------------------------------------------------------


def test_coloracao_amarela_divergente_e_baixa_confianca(tmp_path: Path) -> None:
    """- Tribunal='TJDFT' no cache vs DataJud='TRT/10' → divergente, amarelo.
    - Relator no 2º grau (baixa confiança) sugerido em campo vazio →
      amarelo (não verde, mesmo sendo vazio_sugerido)."""
    out = tmp_path / "test.xlsx"
    schema = _schema_processos_minimo()

    propriedades = {r.nome_notion: None for r in REGRAS_ORIGEM}
    propriedades["Tribunal"] = "TRT/10"
    propriedades["Relator no 2º grau"] = "Maria das Couves"
    res = _resultado(propriedades=propriedades)

    processos_cache: dict[str, dict[str, Any]] = {
        "abc123": {"tribunal": "TJDFT", "relator_2g": None},
    }

    gerar_xlsx([res], processos_cache=processos_cache, schema=schema, output_path=out)

    wb = load_workbook(out)
    ws = wb[SHEET_NAME_DATAJUD]

    # Tribunal: divergente → amarelo
    col_trib: int | None = None
    col_relator: int | None = None
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=1, column=c).value
        if v == "Tribunal":
            col_trib = c
        if v == "Relator no 2º grau":
            col_relator = c

    assert col_trib is not None
    assert col_relator is not None

    fill_trib = ws.cell(row=2, column=col_trib).fill
    assert fill_trib.fgColor.value.upper().endswith(
        CELL_FILL_DIVERGENTE_OU_BAIXA_HEX,
    )
    fill_rel = ws.cell(row=2, column=col_relator).fill
    assert fill_rel.fgColor.value.upper().endswith(
        CELL_FILL_DIVERGENTE_OU_BAIXA_HEX,
    )


def test_celula_igual_nao_e_escrita(tmp_path: Path) -> None:
    """Quando DataJud=Notion, a célula importável fica intacta (não
    escrita) — evita poluir importação com sugestões redundantes."""
    out = tmp_path / "test.xlsx"
    schema = _schema_processos_minimo()
    propriedades = {r.nome_notion: None for r in REGRAS_ORIGEM}
    propriedades["Tribunal"] = "TRT/10"
    res = _resultado(propriedades=propriedades)
    processos_cache: dict[str, dict[str, Any]] = {"abc123": {"tribunal": "TRT/10"}}
    gerar_xlsx([res], processos_cache=processos_cache, schema=schema, output_path=out)

    wb = load_workbook(out)
    ws = wb[SHEET_NAME_DATAJUD]
    col_trib: int | None = None
    col_status: int | None = None
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=1, column=c).value
        if v == "Tribunal":
            col_trib = c
        if v == "Tribunal ▸ status":
            col_status = c
    assert col_trib is not None
    assert col_status is not None
    # Tribunal importável NÃO escrita (igual ao Notion)
    assert ws.cell(row=2, column=col_trib).value is None
    # Status escrita com label canônico
    assert ws.cell(row=2, column=col_status).value == STATUS_DIVERGENCIA["igual"]


# ---------------------------------------------------------------------------
# 7) Validação de lista nos selects fechados
# ---------------------------------------------------------------------------


def test_validacao_lista_nos_selects_fechados(tmp_path: Path) -> None:
    """4 DataValidations (Tribunal, Instância, Status, Fase) apontando
    para a aba _vocabularios."""
    out = tmp_path / "test.xlsx"
    schema = _schema_processos_minimo()
    gerar_xlsx([_resultado()], processos_cache={}, schema=schema, output_path=out)

    wb = load_workbook(out)
    ws = wb[SHEET_NAME_DATAJUD]
    dvs = list(ws.data_validations.dataValidation)
    assert len(dvs) == 4
    formulas = [str(dv.formula1) for dv in dvs]
    for f in formulas:
        assert SHEET_NAME_VOCABULARIOS in f or "_vocabularios" in f


# ---------------------------------------------------------------------------
# 8) Freeze panes e auto_filter
# ---------------------------------------------------------------------------


def test_freeze_panes_e_auto_filter(tmp_path: Path) -> None:
    """freeze_panes='C2' (page_id + Diagnóstico + cabeçalho); auto_filter cobre toda a aba."""
    out = tmp_path / "test.xlsx"
    schema = _schema_processos_minimo()
    gerar_xlsx([_resultado()], processos_cache={}, schema=schema, output_path=out)

    wb = load_workbook(out)
    ws = wb[SHEET_NAME_DATAJUD]
    assert ws.freeze_panes == "C2"
    assert ws.auto_filter.ref is not None
    assert "A1" in ws.auto_filter.ref


# ---------------------------------------------------------------------------
# 9) Meta column contém JSON com hash + fontes + ts
# ---------------------------------------------------------------------------


def test_meta_column_contem_hash_fontes_e_ts(tmp_path: Path) -> None:
    """A última coluna __datajud_meta contém JSON com 3 chaves:
    ts_consulta, fontes_tribunal, hash_payload."""
    out = tmp_path / "test.xlsx"
    schema = _schema_processos_minimo()
    res = _resultado(
        fontes=["trt10", "tst"],
        movs_brutos={"G1": [{"codigo": 26}], "GS": [{"codigo": 51}]},
    )
    gerar_xlsx([res], processos_cache={}, schema=schema, output_path=out)

    wb = load_workbook(out)
    ws = wb[SHEET_NAME_DATAJUD]
    last_col = ws.max_column
    assert ws.cell(row=1, column=last_col).value == "__datajud_meta"
    raw = ws.cell(row=2, column=last_col).value
    assert isinstance(raw, str)
    parsed = json.loads(raw)
    assert "ts_consulta" in parsed
    assert "fontes_tribunal" in parsed
    assert "hash_payload" in parsed
    assert parsed["fontes_tribunal"] == ["trt10", "tst"]
    assert isinstance(parsed["hash_payload"], str) and len(parsed["hash_payload"]) >= 8


# ---------------------------------------------------------------------------
# 10) Header com cor amarelo escuro nos relatores (baixa confiança)
# ---------------------------------------------------------------------------


def test_header_relator_2g_e_stj_amarelo_escuro(tmp_path: Path) -> None:
    """Headers das colunas importáveis 'Relator no 2º grau' e
    'Relator no STJ/TST' usam fill amarelo escuro BF8F00; outros
    headers ficam azul 1F4E79."""
    out = tmp_path / "test.xlsx"
    schema = _schema_processos_minimo()
    gerar_xlsx([_resultado()], processos_cache={}, schema=schema, output_path=out)

    wb = load_workbook(out)
    ws = wb[SHEET_NAME_DATAJUD]
    encontrei_amarelo = 0
    for c in range(1, ws.max_column + 1):
        h = ws.cell(row=1, column=c).value
        fill = ws.cell(row=1, column=c).fill
        cor_hex = (fill.fgColor.value or "").upper() if fill else ""
        if h in ("Relator no 2º grau", "Relator no STJ/TST"):
            assert cor_hex.endswith(HEADER_FILL_LOW_HEX), (
                f"header {h!r}: cor inesperada {cor_hex}"
            )
            encontrei_amarelo += 1
        elif h == "page_id":
            assert cor_hex.endswith(HEADER_FILL_HIGH_HEX)
    assert encontrei_amarelo == 2


# ---------------------------------------------------------------------------
# 11) status_divergencia helper — cobertura dos 5 cenários
# ---------------------------------------------------------------------------


def test_status_divergencia_helper() -> None:
    """status_divergencia cobre os 5 cenários canônicos."""
    # Não enriquecida → nao_coberto
    assert status_divergencia(None, None, enriquecida=False) == "nao_coberto"
    assert status_divergencia("foo", "bar", enriquecida=False) == "nao_coberto"

    # Enriquecidas:
    assert status_divergencia(None, "TRT/10", enriquecida=True) == "vazio_sugerido"
    assert status_divergencia("", "TRT/10", enriquecida=True)   == "vazio_sugerido"
    assert status_divergencia("TRT/10", None, enriquecida=True) == "datajud_vazio"
    assert status_divergencia(None, None, enriquecida=True)     == "datajud_vazio"
    assert status_divergencia("TRT/10", "TRT/10", enriquecida=True) == "igual"
    assert status_divergencia("TJDFT", "TRT/10", enriquecida=True) == "divergente"
    # Lista vs string
    assert status_divergencia(["A", "B"], "A, B", enriquecida=True) == "divergente"


# ---------------------------------------------------------------------------
# 12) Múltiplas linhas
# ---------------------------------------------------------------------------


def test_multiplas_linhas_e_diagnosticos_diferentes(tmp_path: Path) -> None:
    """3 resultados (OK, Não encontrado, Erro) ocupam linhas 2-4."""
    out = tmp_path / "test.xlsx"
    schema = _schema_processos_minimo()
    propriedades_ok = {r.nome_notion: None for r in REGRAS_ORIGEM}
    propriedades_ok["Tribunal"] = "TRT/10"

    resultados = [
        _resultado(page_id="p1", diagnostico="OK", propriedades=propriedades_ok),
        _resultado(page_id="p2", diagnostico="Não encontrado"),
        _resultado(page_id="p3", diagnostico="Erro: HTTP 503"),
    ]
    gerar_xlsx(resultados, processos_cache={}, schema=schema, output_path=out)

    wb = load_workbook(out)
    ws = wb[SHEET_NAME_DATAJUD]
    assert ws.cell(row=2, column=1).value == "p1"
    assert ws.cell(row=3, column=1).value == "p2"
    assert ws.cell(row=4, column=1).value == "p3"
    assert ws.cell(row=2, column=2).value == "OK"
    assert ws.cell(row=3, column=2).value == "Não encontrado"
    assert ws.cell(row=4, column=2).value == "Erro: HTTP 503"
